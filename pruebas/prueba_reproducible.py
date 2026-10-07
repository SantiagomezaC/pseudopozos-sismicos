# -*- coding: utf-8 -*-
"""
Prueba reproducible de pseudopozos_sismicos.py
===============================================
Verifica, sin interfaz gráfica y con datos sintéticos generados aquí mismo, los componentes
de los que dependen los resultados:

  1. Ecuación de espesor e = V * ΔTWT / 2, contrastada con las cinco primeras filas de la
     base de puntos de control del proyecto (línea CV-1989-950, V_Q = 1500 m/s).
  2. Tratamiento de capas ausentes y de secciones en profundidad (columnas "NA").
  3. Identificación de la línea en el catálogo a partir de textos leídos por OCR, incluidos
     errores típicos y nombres parecidos que no deben aceptarse.
  4. Exactitud geoespacial: los puntos quedan sobre la geometría de la línea BIP.
  5. Exportación a Excel y CSV.
  6. OCR de Windows sobre una imagen generada (se omite si el equipo no lo tiene).

Uso:  python pruebas/prueba_reproducible.py      (código de salida 0 si todo es correcto)
"""
import os
import sys
import tempfile

import numpy as np
import pyogrio
import shapely
from shapely.geometry import LineString

# El programa está en la carpeta superior; se importa como módulo (no abre la ventana).
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import pseudopozos_sismicos as pp

resultados = []                                   # (nombre de la prueba, aprobada, detalle)


def verificar(nombre, condicion, detalle=""):
    """Registra e imprime el resultado de una comprobación."""
    resultados.append((nombre, bool(condicion), detalle))
    print(f"[{'OK ' if condicion else 'FALLA'}] {nombre}" + (f"  ({detalle})" if detalle else ""))


tmp = tempfile.mkdtemp(prefix="pseudopozos_prueba_")   # carpeta temporal de trabajo

# -----------------------------------------------------------------------------------------
# 1. Ecuación de espesor contra la tabla de puntos de control del proyecto
# -----------------------------------------------------------------------------------------
# Filas 1 a 5 de la base original: (TWT base Q, TWT base N, espesor total esperado en m)
tabla = [(0.47, 0.626, 542.6952), (0.47, 0.624, 540.2568), (0.50, 0.625, 527.4),
         (0.52, 0.620, 511.92), (0.50, 0.624, 526.1808)]
# Calibración sintética: 1 píxel vertical = 0,001 s TWT y 100 píxeles horizontales por km.
xs = [100.0 * i for i in range(len(tabla))]
cfg = dict(linea="CV-1989-950", modo="TWT", xcal=[0.0, 400.0], ycal=[0.0, 1000.0], ref=[0.0, 1.0],
           long_km=4.0, ajustar_bip=True, offset_km=0.0, capas=["Q", "N"],
           vel={"Q": 1500.0, "N": 2438.4, "P": 2438.4},           # la tabla original usó 1500 m/s
           horizontes={"sup": [],
                       "Q": [[x, q * 1000] for x, (q, _, _) in zip(xs, tabla)],
                       "bas": [[x, n * 1000] for x, (_, n, _) in zip(xs, tabla)]},
           pozos_km=[0.0, 1.0, 2.0, 3.0, 4.0], punto_inicial=1)
linea_recta = pp.LineaGeo([(1000000.0, 1000000.0), (1004000.0, 1000000.0)], "EPSG:3116")
filas, _ = pp.calcular(cfg, linea_recta)
obtenidos = [round(f["esp_total"], 4) for f in filas]
esperados = [t[2] for t in tabla]
verificar("Espesor total = tabla del proyecto (filas 1-5)", np.allclose(obtenidos, esperados, atol=1e-4),
          f"obtenido {obtenidos}")
verificar("Espesor Q fila 1 = 1500 x 0,47 / 2 = 352,5 m", abs(filas[0]["esp"]["Q"] - 352.5) < 1e-9)
verificar("Espesor N fila 1 = 2438,4 x 0,156 / 2 = 190,1952 m", abs(filas[0]["esp"]["N"] - 190.1952) < 1e-9)

# Con la velocidad del Cuaternario adoptada en el proyecto (1150 m/s)
cfg_1150 = dict(cfg, vel={"Q": 1150.0, "N": 2438.4, "P": 2438.4})
f1150, _ = pp.calcular(cfg_1150, linea_recta)
verificar("Espesor Q con 1150 m/s = 270,25 m", abs(f1150[0]["esp"]["Q"] - 270.25) < 1e-9)

# -----------------------------------------------------------------------------------------
# 2. Capas ausentes y sección en profundidad
# -----------------------------------------------------------------------------------------
verificar("Paleógeno ausente en la línea -> espesor 0 y TWT vacío",
          all(f["esp"]["P"] == 0 and f["twt"]["P"] is None for f in filas))
cfg_acu = dict(cfg, horizontes=dict(cfg["horizontes"], Q=[[200.0, 500.0], [400.0, 500.0]]))
f_acu, _ = pp.calcular(cfg_acu, linea_recta)        # Q solo existe desde el km 2
verificar("Capa acuñada (sin horizonte en esa posición) -> espesor 0",
          f_acu[0]["esp"]["Q"] == 0 and f_acu[1]["esp"]["Q"] == 0 and f_acu[2]["esp"]["Q"] > 0)
f_prof, _ = pp.calcular(dict(cfg, modo="PROF"), linea_recta)
fila_prof = pp.fila_principal(f_prof[0])
verificar("Sección en profundidad: columnas por capa = 'NA' y espesor total en m",
          fila_prof[5:14] == ["NA"] * 9 and abs(f_prof[0]["esp_total"] - 626.0) < 1e-9)

# -----------------------------------------------------------------------------------------
# 3. Identificación de la línea a partir de textos leídos por OCR
# -----------------------------------------------------------------------------------------
x0, y0 = 1000000.0, 1100000.0
lineas = {   # catálogo sintético en EPSG:3116, con nombres deliberadamente parecidos
    "CV-1979-08": [LineString([(x0 + 20000, y0 + 21000), (x0 + 41000, y0 + 40000)]),
                   LineString([(x0, y0), (x0 + 8000, y0 + 9500), (x0 + 20000, y0 + 21000)])],
    "CV-1979-080": [LineString([(x0, y0 + 50000), (x0 + 9000, y0 + 50000)])],
    "CV-1979-8A": [LineString([(x0, y0 + 60000), (x0 + 9000, y0 + 60000)])],
    "CV-1989-950": [LineString([(x0, y0 + 5000), (x0 + 10000, y0 + 5000)])],
}
geoms, nombres, programas = [], [], []
for n, partes in lineas.items():
    for g in partes:
        geoms.append(g)
        nombres.append(n)
        programas.append(n.rsplit("-", 1)[0])           # p. ej. "CV-1979" como programa
ruta_cat = os.path.join(tmp, "catalogo.gpkg")
pyogrio.raw.write(ruta_cat, np.array(shapely.to_wkb(geoms), dtype=object),
                  [np.array(nombres, dtype=object), np.array(programas, dtype=object)],
                  ["SURVEY_NAM", "PROGRAMA_L"], layer="sismica", driver="GPKG",
                  geometry_type="LineString", crs="EPSG:3116")
cat = pp.BaseBIP(ruta_cat, "sismica")
casos = {"CVI 979-08": "CV-1979-08",                 # I leída en lugar de 1, espacio intermedio
         "Line cv-1979-O8 migrated": "CV-1979-08",   # O en lugar de 0, texto alrededor
         "Seccion_CV_1979_08": "CV-1979-08",         # nombre de archivo
         "CV-1979-080": "CV-1979-080"}               # no debe confundirse con CV-1979-08
for texto, esperado in casos.items():
    r = cat.identificar([texto])
    verificar(f"Identificación de «{texto}»", len(r) == 1 and r[0][1] == esperado,
              f"resultado {[v for _, v, _ in r]}")
verificar("Texto sin nombre de línea no identifica nada", cat.identificar(["Escala 2 km 1979"]) == [])

# -----------------------------------------------------------------------------------------
# 4. Exactitud geoespacial sobre la línea BIP
# -----------------------------------------------------------------------------------------
xy, nota = cat.geometria("SURVEY_NAM", "CV-1979-08")
verificar("Los dos tramos de la línea se unen en una polilínea continua", len(xy) == 4, nota)
L = pp.LineaGeo(xy, cat.crs)
verificar("Rumbo de la línea SW -> NE", L.rumbo() == ("SW", "NE"))
verificar("Orientación: extremo izquierdo SW = primer vértice",
          np.allclose(L.orientada("SW").xy[0], xy[0]) and np.allclose(L.orientada("NE").xy[0], xy[-1]))
puntos = [L.punto(fr) for fr in np.linspace(0, 1, 34)]
desv = max(p["desv_m"] for p in puntos)
verificar("Desviación máxima de 34 puntos respecto a la línea < 1 micra", desv < 1e-6, f"{desv:.2e} m")
verificar("Extremos coinciden con los vértices BIP",
          np.allclose([puntos[0]["x"], puntos[0]["y"]], xy[0]) and
          np.allclose([puntos[-1]["x"], puntos[-1]["y"]], xy[-1]))
p = L.punto(0.1234)
largo = float(np.sum(np.hypot(*np.diff(xy, axis=0).T)))
recorrido = L.geom.project(shapely.Point(p["x"], p["y"]))
verificar("Distancia recorrida = fracción x longitud", abs(recorrido - 0.1234 * largo) < 1e-6,
          f"{recorrido:.4f} m")

# -----------------------------------------------------------------------------------------
# 5. Exportación
# -----------------------------------------------------------------------------------------
ruta_xlsx = os.path.join(tmp, "pseudopozos.xlsx")
ruta_csv = pp.exportar_excel(filas, {"Línea": "CV-1989-950"}, ruta_xlsx)
pp.exportar_excel(f_prof, {"Línea": "CV-1989-950"}, ruta_xlsx, agregar=True)
from openpyxl import load_workbook
wb = load_workbook(ruta_xlsx)
verificar("Excel con hojas Pseudopozos, GIS y Metadatos",
          wb.sheetnames == ["Pseudopozos", "GIS", "Metadatos"])
verificar("Agregar a la base existente conserva las filas anteriores",
          wb["Pseudopozos"].max_row == 1 + 2 * len(filas))
encabezado_csv = open(ruta_csv, encoding="utf-8-sig").readline().strip().split(",")
verificar("CSV con columnas Longitud/Latitud para SIG", encabezado_csv == pp.ENC_GIS)

# -----------------------------------------------------------------------------------------
# 6. OCR de Windows (opcional)
# -----------------------------------------------------------------------------------------
from PIL import Image, ImageDraw, ImageFont
ruta_img = os.path.join(tmp, "rotulo.png")
im = Image.new("RGB", (700, 120), "white")
ImageDraw.Draw(im).text((20, 30), "CV-1979-08", fill="black", font=ImageFont.load_default(48))
im.save(ruta_img)
ocr = pp.ocr_imagen(ruta_img)
if ocr is None:
    print("[--] OCR de Windows no disponible en este equipo: prueba omitida")
else:
    r = cat.identificar([t for t, _ in ocr])
    verificar("OCR + catálogo identifican CV-1979-08 en una imagen", r and r[0][1] == "CV-1979-08",
              f"leído {[t for t, _ in ocr][:3]}")

# -----------------------------------------------------------------------------------------
fallas = [r for r in resultados if not r[1]]
print(f"\n{len(resultados) - len(fallas)} de {len(resultados)} comprobaciones correctas.")
sys.exit(1 if fallas else 0)
