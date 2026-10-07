# -*- coding: utf-8 -*-
"""
==========================================================================================
PSEUDOPOZOS SÍSMICOS
Generación de pseudopozos georreferenciados a partir de secciones sísmicas interpretadas
==========================================================================================

Propósito
---------
El programa convierte una sección sísmica 2D interpretada (una imagen) en una base de datos
de "pseudopozos": puntos de control distribuidos a lo largo de la línea sísmica, cada uno
con coordenadas geográficas y con el espesor de las unidades estratigráficas comprendidas
entre la superficie y el límite con el basamento. La base resultante se exporta a Excel y
CSV en un formato que QGIS y ArcGIS cargan directamente como capa de puntos.

Flujo metodológico
------------------
1. Identificación de la línea. El texto impreso en la imagen (leído con el OCR integrado de
   Windows) y el nombre del archivo se comparan con el catálogo de navegación sísmica 2D del
   Banco de Información Petrolera (BIP) de la Agencia Nacional de Hidrocarburos (ANH). Solo
   se acepta un nombre que exista en el catálogo.
2. Calibración de la imagen. El usuario marca los extremos de la sección (eje horizontal) y
   dos niveles de referencia de tiempo doble (TWT) o de profundidad (eje vertical). Con ello
   se establece una transformación lineal píxel -> (distancia, TWT o profundidad).
3. Digitalización. El usuario traza la superficie (opcional), la base de cada unidad y el
   límite con el basamento.
4. Cálculo. En cada pseudopozo se leen los horizontes y se calcula el espesor de cada unidad
   con su velocidad interválica:

            e_i = V_i * (TWT_base,i - TWT_tope,i) / 2

   donde la división por 2 convierte el tiempo doble (ida y vuelta) en tiempo sencillo.
   Convención adoptada en el proyecto:
        Capa 1 = Cuaternario (Q)   V = 1150   m/s
        Capa 2 = Neógeno     (N)   V = 2438.4 m/s
        Capa 3 = Paleógeno   (P)   V = 2438.4 m/s
   En secciones convertidas a profundidad (km) el espesor total es la diferencia directa de
   profundidades y las columnas por capa se reportan como "NA" (no aplica).
5. Georreferenciación. Cada pseudopozo se sitúa sobre la geometría oficial de la línea BIP,
   interpolando entre sus vértices en el sistema de coordenadas original (MAGNA-SIRGAS
   Bogotá, EPSG:3116, en el catálogo de la ANH). Así el punto queda exactamente sobre la
   línea; la desviación de cada punto respecto a ella se calcula y se reporta.

Reproducibilidad
----------------
Requiere Windows 10/11 (por el OCR) y Python 3.11 o superior. Las versiones exactas de las
librerías están fijadas en requirements.txt; instalar.bat crea el entorno y ejecutar.bat
abre el programa. La prueba pruebas/prueba_reproducible.py verifica los cálculos sin
interfaz gráfica.

Uso:   python pseudopozos_sismicos.py
"""

# =========================================================================================
# 1. LIBRERÍAS
# =========================================================================================
# --- Biblioteca estándar de Python (no requiere instalación) ---
import asyncio            # ejecuta las llamadas asíncronas del motor OCR de Windows
import csv                # escritura del archivo CSV para los SIG
import datetime           # fecha de exportación y conversión de fechas del servicio ANH
import difflib            # sugerencias de nombres parecidos en la búsqueda manual
import json               # lectura/escritura de la configuración y de los proyectos
import math               # funciones trigonométricas para rumbos y direcciones
import os                 # manejo de rutas y archivos
import re                 # expresiones regulares para normalizar nombres de líneas
import sys                # argumentos de línea de comandos (modo de autoprueba)
import threading          # descarga del catálogo ANH sin congelar la interfaz
import traceback          # mensajes de error detallados para el usuario
import unicodedata        # eliminación de tildes al comparar textos
import urllib.error       # errores de red al consultar el servicio de la ANH
import urllib.parse       # construcción de las consultas al servicio ArcGIS REST
import urllib.request     # peticiones HTTP al servicio de la ANH
import tkinter as tk      # interfaz gráfica de escritorio
from tkinter import ttk, filedialog, messagebox, simpledialog   # controles y diálogos

# --- Librerías científicas y geoespaciales (ver requirements.txt) ---
import numpy as np                     # cálculo numérico vectorizado e interpolación
from PIL import Image                  # lectura de la imagen de la sección sísmica
import matplotlib                      # visualización de la imagen y de lo digitalizado
matplotlib.use("TkAgg")                # se integra matplotlib dentro de la ventana Tkinter
from matplotlib.figure import Figure    # lienzo donde se dibuja la sección
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg, NavigationToolbar2Tk
import pyogrio                         # lectura/escritura de formatos SIG (shp, gpkg, gdb...)
import shapely                         # geometría vectorial (líneas, puntos, distancias)
from shapely.geometry import MultiLineString
from shapely.ops import linemerge      # unión de tramos contiguos de una misma línea
from pyproj import CRS, Geod, Transformer   # sistemas de referencia y transformaciones
from openpyxl import Workbook, load_workbook # creación y edición de archivos Excel
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

Image.MAX_IMAGE_PIXELS = None   # se desactiva el límite de tamaño de PIL: las secciones
                                # sísmicas escaneadas suelen superar los 89 millones de píxeles

# =========================================================================================
# 2. PARÁMETROS GENERALES
# =========================================================================================
APP = "Pseudopozos sísmicos"   # nombre que aparece en ventanas y mensajes
VERSION = "2.0"                # versión del programa, registrada en los proyectos guardados

# Servicio ArcGIS REST del geovisor de la ANH; la capa 2 corresponde a "Sísmica 2D".
URL_ANH = ("https://geovisor.anh.gov.co/server/rest/services/GEOVISOR_v32/"
           "ANH_InsGDB/MapServer/2")

# Carpeta de trabajo del programa en el perfil del usuario (%APPDATA%\Pseudopozos): allí se
# guardan la configuración y la copia local del catálogo BIP descargado.
DIR_APP = os.path.join(os.environ.get("APPDATA", os.path.expanduser("~")), "Pseudopozos")
RUTA_CFG = os.path.join(DIR_APP, "config.json")                     # configuración persistente
RUTA_CATALOGO_ANH = os.path.join(DIR_APP, "sismica_2d_anh.gpkg")    # catálogo en GeoPackage

# Unidades estratigráficas. Para cada clave se guarda:
# (número de capa, nombre, velocidad interválica por defecto en m/s, color de dibujo).
CAPAS = {
    "Q": (1, "Cuaternario", 1150.0, "#8c8c8c"),
    "N": (2, "Neógeno", 2438.4, "#e0b800"),
    "P": (3, "Paleógeno", 2438.4, "#f07c1e"),
}
ORDEN_CAPAS = ["Q", "N", "P"]   # orden estratigráfico de techo a base

COLOR_SUP = "#00bcd4"           # color del horizonte de superficie
COLOR_BAS = "#1f4fff"           # color del límite con el basamento
COLOR_POZO = "#d00000"          # color de las verticales de los pseudopozos
SNAP_PX = 10       # radio (píxeles de pantalla) dentro del cual un clic se "pega" a otro
                   # horizonte; garantiza que los acuñamientos coincidan exactamente
TOL_PX = 2.0       # diferencias menores a 2 píxeles de imagen se consideran espesor nulo,
                   # para no reportar espesores ficticios producto del pulso de la mano

# Opciones para indicar qué extremo de la línea BIP está a la izquierda de la sección.
EXTREMOS = ["Primer vértice BIP", "Último vértice BIP", "W", "SW", "S", "NW", "N", "NE", "E", "SE"]

# Rótulos de orientación que pueden aparecer en una sección (en inglés y en español: O = oeste).
RE_CARDINAL = re.compile(r"^(N|S|E|W|O|NE|NW|SE|SW|NO|SO|NNE|ENE|ESE|SSE|SSW|WSW|WNW|NNW|"
                         r"SSO|OSO|ONO|NNO)$")


# =========================================================================================
# 3. CONFIGURACIÓN PERSISTENTE
#    Permite que el catálogo BIP se configure una sola vez y se cargue automáticamente.
# =========================================================================================
def leer_cfg():
    """Lee el archivo de configuración; si no existe o está dañado devuelve un diccionario vacío."""
    try:
        with open(RUTA_CFG, encoding="utf-8") as fh:   # se abre en UTF-8 por las tildes
            return json.load(fh)
    except (OSError, ValueError):                      # archivo ausente o JSON inválido
        return {}


def guardar_cfg(**cambios):
    """Actualiza solo las claves indicadas y conserva el resto de la configuración."""
    cfg = leer_cfg()                                   # configuración vigente
    cfg.update(cambios)                                # se sobrescriben las claves nuevas
    os.makedirs(DIR_APP, exist_ok=True)                # se crea la carpeta si no existe
    with open(RUTA_CFG, "w", encoding="utf-8") as fh:
        json.dump(cfg, fh, ensure_ascii=False, indent=1)


# =========================================================================================
# 4. UTILIDADES DE NOMBRES Y GEOMETRÍA
# =========================================================================================
def _sin_tildes(s):
    """Elimina tildes y diacríticos ('Neógeno' -> 'Neogeno') para comparar textos."""
    return unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()


def tokens_nombre(s, ocr=False):
    """Descompone el nombre de una línea en bloques de letras y de números.

    'CV-1979-08', 'cv 1979 8' y 'CV_1979_08' producen los mismos bloques ('CV','1979','8'):
    se ignoran mayúsculas, separadores y ceros a la izquierda. Con ocr=True se corrigen las
    confusiones típicas del OCR junto a dígitos (I, L, | o ! leídos en lugar de 1; O en lugar
    de 0). La corrección se aplica por igual al texto leído y al catálogo, de modo que la
    comparación es simétrica y no puede crear coincidencias que no existan en el BIP.
    """
    s = _sin_tildes(s).upper()                                  # sin tildes y en mayúsculas
    if ocr:
        s = re.sub(r"(?<=\d)[IL|!]|[IL|!](?=\d)", "1", s)       # I/L/|/! junto a dígito -> 1
        s = re.sub(r"(?<=\d)O|O(?=\d)", "0", s)                 # O junto a dígito -> 0
    toks = re.findall(r"[A-Z]+|\d+", s)                         # bloques de letras o dígitos
    return tuple((t.lstrip("0") or "0") if t.isdigit() else t for t in toks)  # '08' -> '8'


def normalizar(nombre):
    """Forma canónica de un nombre: 'CV-1979-08', 'cv 1979 8' y 'CV_1979_08' -> 'CV-1979-8'."""
    return "-".join(tokens_nombre(nombre))


def vector_rumbo(etiqueta):
    """Convierte un rumbo escrito ('SW', 'NE', 'O'...) en un vector unitario (este, norte)."""
    v = {"N": (0, 1), "S": (0, -1), "E": (1, 0), "W": (-1, 0), "O": (-1, 0)}  # O = oeste
    x = sum(v[c][0] for c in etiqueta)      # componente este: suma de cada letra
    y = sum(v[c][1] for c in etiqueta)      # componente norte
    n = math.hypot(x, y) or 1.0             # norma del vector (se evita dividir por cero)
    return x / n, y / n


def rumbo_8(dx, dy):
    """Clasifica una dirección (dx hacia el este, dy hacia el norte) en uno de 8 rumbos."""
    ang = (math.degrees(math.atan2(dx, dy)) + 360) % 360       # azimut 0-360° desde el norte
    return ["N", "NE", "E", "SE", "S", "SW", "W", "NW"][int((ang + 22.5) // 45) % 8]


def encadenar(partes):
    """Une en una sola polilínea ordenada los tramos de una misma línea sísmica.

    En el catálogo una línea puede estar partida en varios registros. Primero se unen los
    tramos que comparten extremos (linemerge); si quedan tramos separados, se encadenan
    uniendo cada vez el extremo más cercano. Devuelve las coordenadas y el número de
    discontinuidades que debieron unirse con un segmento recto (se informa al usuario).
    """
    lineas = [np.asarray(p.coords)[:, :2] for p in partes if len(p.coords) >= 2]  # solo X,Y
    unida = linemerge(MultiLineString(lineas))                  # une tramos que se tocan
    if unida.geom_type == "LineString":                         # quedó una sola línea continua
        return np.asarray(unida.coords)[:, :2], 0
    tramos = [np.asarray(g.coords)[:, :2] for g in unida.geoms] # tramos aún separados
    tramos.sort(key=lambda s: -np.sum(np.hypot(*np.diff(s, axis=0).T)))  # el más largo primero
    cadena = tramos.pop(0)                                      # se parte del tramo más largo
    huecos = 0
    while tramos:                                               # mientras queden tramos sueltos
        mejor = None
        for j, s in enumerate(tramos):
            for inv in (False, True):                           # se prueba en ambos sentidos
                ss = s[::-1] if inv else s
                d_fin = np.hypot(*(cadena[-1] - ss[0]))         # distancia al final de la cadena
                d_ini = np.hypot(*(ss[-1] - cadena[0]))         # distancia al inicio de la cadena
                for d, lado in ((d_fin, "fin"), (d_ini, "ini")):
                    if mejor is None or d < mejor[0]:           # se conserva la unión más corta
                        mejor = (d, j, ss, lado)
        _, j, ss, lado = mejor
        tramos.pop(j)                                           # el tramo elegido se retira
        cadena = np.vstack([cadena, ss]) if lado == "fin" else np.vstack([ss, cadena])
        huecos += 1                                             # se contabiliza la discontinuidad
    return cadena, huecos


def evaluar(puntos, x):
    """Valor vertical (en píxeles) de un horizonte digitalizado en la abscisa x.

    El horizonte se trata como una función de x interpolada linealmente entre sus vértices.
    Fuera de su extensión devuelve NaN: así se interpreta que la unidad está ausente en esa
    posición (por ejemplo, donde una capa se acuña).
    """
    if len(puntos) < 2:                                    # con menos de 2 vértices no hay línea
        return float("nan")
    P = np.asarray(sorted(puntos), dtype=float)            # vértices ordenados de izquierda a derecha
    return float(np.interp(x, P[:, 0], P[:, 1], left=np.nan, right=np.nan))


# =========================================================================================
# 5. GEOMETRÍA DE LA LÍNEA SÍSMICA (exactitud geoespacial)
# =========================================================================================
class LineaGeo:
    """Línea sísmica representada en el sistema de coordenadas ORIGINAL del BIP.

    Los pseudopozos se interpolan linealmente entre los vértices de la geometría tal como
    está almacenada, es decir, sobre los mismos segmentos rectos que dibuja cualquier SIG.
    Por construcción el punto cae sobre la línea; solo después se transforma a WGS84
    (EPSG:4326) y a MAGNA-SIRGAS Origen Nacional (EPSG:9377). Si se reproyectara la línea
    antes de interpolar, los segmentos rectos en un sistema dejarían de serlo en el otro y
    los puntos se separarían ligeramente de la geometría oficial.

    La distancia a lo largo de la línea se mide en metros: plana si el sistema es proyectado
    (como EPSG:3116) y geodésica sobre el elipsoide WGS84 si es geográfico.
    """
    geod = Geod(ellps="WGS84")    # calculadora geodésica, usada solo con coordenadas geográficas

    def __init__(self, xy, crs):
        self.crs = CRS.from_user_input(crs)                 # sistema de referencia del BIP
        xy = np.asarray(xy, float)[:, :2]                   # coordenadas X,Y (se descarta Z)
        if len(xy) >= 2:                                    # se eliminan vértices repetidos
            keep = np.concatenate([[True], np.hypot(*np.diff(xy, axis=0).T) > 0])
            xy = xy[keep]
        if len(xy) < 2:
            raise ValueError("La geometría de la línea tiene menos de dos vértices distintos.")
        self.xy = xy                                        # vértices en el sistema original
        # Transformaciones del sistema original a WGS84 y a Origen Nacional (always_xy=True
        # fija el orden (x, y) = (este, norte) / (longitud, latitud)).
        self._a4326 = Transformer.from_crs(self.crs, CRS.from_epsg(4326), always_xy=True)
        self._a9377 = Transformer.from_crs(self.crs, CRS.from_epsg(9377), always_xy=True)
        lon, lat = self._a4326.transform(xy[:, 0], xy[:, 1])   # vértices en grados (para rumbos)
        self.lon, self.lat = np.asarray(lon), np.asarray(lat)
        self.proyectado = self.crs.is_projected             # ¿coordenadas planas en metros?
        if self.proyectado:
            # factor de la unidad del sistema a metros (1.0 si ya está en metros)
            self.factor = self.crs.axis_info[0].unit_conversion_factor or 1.0
            seg = np.hypot(*np.diff(xy, axis=0).T) * self.factor   # longitud de cada segmento
        else:
            self.factor = 111319.49   # metros por grado; solo para expresar desviaciones en m
            _, _, seg = self.geod.inv(self.lon[:-1], self.lat[:-1], self.lon[1:], self.lat[1:])
        self.cum = np.concatenate([[0.0], np.cumsum(seg)])  # distancia acumulada en cada vértice
        self.geom = shapely.LineString(xy)                  # geometría para verificar desviaciones

    @property
    def epsg(self):
        """Código EPSG del sistema original (o su nombre si no tiene código)."""
        e = self.crs.to_epsg()
        return f"EPSG:{e}" if e else self.crs.name

    @property
    def longitud_km(self):
        """Longitud total de la línea en kilómetros."""
        return self.cum[-1] / 1000.0

    def invertida(self):
        """La misma línea recorrida en sentido contrario (último vértice primero)."""
        return LineaGeo(self.xy[::-1], self.crs)

    def _direccion(self):
        """Vector aproximado primer -> último vértice (dx corregido por la convergencia de meridianos)."""
        dx = (self.lon[-1] - self.lon[0]) * math.cos(math.radians(self.lat.mean()))
        dy = self.lat[-1] - self.lat[0]
        return dx, dy

    def rumbo(self):
        """Rumbos (8 direcciones) hacia los que se encuentran el primer y el último vértice."""
        dx, dy = self._direccion()
        return rumbo_8(-dx, -dy), rumbo_8(dx, dy)

    def coseno_primer_vertice(self, etiqueta):
        """Coseno del ángulo entre el rumbo 'etiqueta' y la dirección último -> primer vértice.

        Positivo: el rótulo apunta hacia el primer vértice; negativo: hacia el último.
        """
        dx, dy = self._direccion()
        vx, vy = vector_rumbo(etiqueta)
        return (-dx * vx - dy * vy) / (math.hypot(dx, dy) or 1.0)

    def orientada(self, extremo_izq):
        """Devuelve la línea con su primer vértice en el extremo IZQUIERDO de la sección."""
        if extremo_izq == "Primer vértice BIP":
            return self
        if extremo_izq == "Último vértice BIP":
            return self.invertida()
        # Si se indica un rumbo (p. ej. 'SW'), se invierte cuando apunta al último vértice.
        return self if self.coseno_primer_vertice(extremo_izq) >= 0 else self.invertida()

    def punto(self, fraccion):
        """Punto situado a la fracción [0, 1] de la longitud, contada desde el primer vértice.

        Devuelve las coordenadas en el sistema original (x, y), en WGS84 (lon, lat), en
        EPSG:9377 (este, norte) y la desviación en metros respecto a la línea, que por
        construcción es nula salvo por el redondeo de la aritmética de punto flotante.
        """
        d = min(max(fraccion, 0.0), 1.0) * self.cum[-1]           # distancia objetivo en m
        i = int(np.searchsorted(self.cum, d, side="right") - 1)   # segmento que la contiene
        i = min(max(i, 0), len(self.cum) - 2)                     # se acota al último segmento
        t = (d - self.cum[i]) / (self.cum[i + 1] - self.cum[i])   # posición relativa (0-1) en él
        x, y = self.xy[i] + t * (self.xy[i + 1] - self.xy[i])     # interpolación lineal
        lon, lat = self._a4326.transform(x, y)                    # coordenadas geográficas
        este, norte = self._a9377.transform(x, y)                 # coordenadas Origen Nacional
        desv = self.geom.distance(shapely.Point(x, y)) * self.factor   # control de calidad
        return dict(x=float(x), y=float(y), lon=float(lon), lat=float(lat),
                    este=float(este), norte=float(norte), desv_m=float(desv))


# =========================================================================================
# 6. CATÁLOGO DE NAVEGACIÓN SÍSMICA DEL BIP
# =========================================================================================
class BaseBIP:
    """Catálogo de líneas sísmicas 2D leído de un archivo SIG (shp, gpkg, gdb, kml...)."""

    # Palabras que suelen aparecer en el nombre del campo que guarda el nombre de la línea.
    PALABRAS = ("LINE_NAME", "LINENAME", "NOMBRE_LIN", "NOM_LINEA", "LINEA", "LINE",
                "NOMBRE", "NAME", "SURVEY", "PROGRAMA")
    # Palabras del campo numérico que ordena los puntos de tiro cuando la navegación es de puntos.
    PALABRAS_SP = ("SHOTPOINT", "SHOT", "SP", "CDP", "TRACE", "TRAZA", "PUNTO", "PT")

    def __init__(self, ruta, capa=None):
        # pyogrio lee de una vez geometrías (en formato WKB) y atributos de la capa.
        meta, _, geom, datos = pyogrio.raw.read(ruta, layer=capa)
        self.ruta, self.capa = ruta, capa
        self.campos = list(meta["fields"])                                 # nombres de campos
        self.datos = {c: datos[i] for i, c in enumerate(self.campos)}      # atributos por campo
        self.geoms = shapely.from_wkb(geom)                                # geometrías shapely
        self.crs = meta.get("crs")                                         # sistema de referencia
        self.campos_texto = [c for c in self.campos if self.datos[c].dtype == object]  # de texto

    @staticmethod
    def capas(ruta):
        """Lista de capas (nombre, tipo de geometría) contenidas en el archivo o la .gdb."""
        return [(str(n), str(t)) for n, t in pyogrio.list_layers(ruta)]

    def campo_sugerido(self):
        """Campo de texto que con mayor probabilidad contiene el nombre de la línea."""
        def puntaje(c):
            cu = c.upper()
            for i, p in enumerate(self.PALABRAS):       # las primeras palabras pesan más
                if p in cu:
                    return len(self.PALABRAS) - i
            return 0
        cand = sorted(self.campos_texto, key=puntaje, reverse=True)
        return cand[0] if cand else (self.campos[0] if self.campos else None)

    def _valores(self, campo):
        """Valores distintos y no vacíos de un campo, como texto."""
        return {str(v) for v in self.datos[campo] if v is not None and str(v).strip()}

    def buscar(self, nombre, campo):
        """Búsqueda manual por nombre. Devuelve (campo, coincidencias_exactas, sugerencias).

        Primero busca coincidencias exactas tras normalizar (en el campo elegido y luego en
        los demás campos de texto); si no las hay, propone nombres que contienen al buscado
        o se le parecen (difflib), para que el usuario decida.
        """
        objetivo = normalizar(nombre)
        for c in [campo] + [c for c in self.campos_texto if c != campo]:
            exactos = sorted(v for v in self._valores(c) if normalizar(v) == objetivo)
            if exactos:
                return c, exactos, []
        norm = {}                                            # forma normalizada -> nombres reales
        for v in self._valores(campo):
            norm.setdefault(normalizar(v), []).append(v)
        contiene = [k for k in norm if objetivo and (objetivo in k or k in objetivo)]
        cerca = difflib.get_close_matches(objetivo, list(norm), n=15, cutoff=0.6)
        sug = []
        for k in contiene + cerca:                           # sugerencias sin repetir
            for v in sorted(norm[k]):
                if v not in sug:
                    sug.append(v)
        return campo, [], sug[:20]

    def geometria(self, campo, valor):
        """Coordenadas ordenadas de la línea (en el sistema original) y una nota descriptiva.

        Admite navegación como líneas (uno o varios registros por línea) o como puntos de
        tiro; en este último caso los puntos se ordenan por el campo de punto de tiro.
        """
        idx = [i for i, v in enumerate(self.datos[campo]) if str(v) == valor]   # registros
        partes, puntos = [], []
        campo_sp = next((c for c in self.campos if c not in self.campos_texto       # campo SP
                         and any(p in c.upper() for p in self.PALABRAS_SP)), None)
        for i in idx:
            g = self.geoms[i]
            if g is None or g.is_empty:                     # registros sin geometría se ignoran
                continue
            t = g.geom_type
            if t == "LineString":
                partes.append(g)
            elif t == "MultiLineString":
                partes.extend(g.geoms)                       # cada parte por separado
            elif t in ("Point", "MultiPoint"):
                orden = float(self.datos[campo_sp][i]) if campo_sp else float(i)
                subs = [g] if t == "Point" else list(g.geoms)
                for j, s in enumerate(subs):
                    puntos.append((orden, j, s.x, s.y))
        if partes:                                           # navegación como líneas
            xy, huecos = encadenar(partes)
            nota = f"{len(idx)} registro(s)"
            if huecos:
                nota += f"; {huecos} discontinuidad(es) unida(s) en línea recta"
            return xy, nota
        if len(puntos) >= 2:                                 # navegación como puntos de tiro
            puntos.sort()
            nota = f"{len(puntos)} puntos de tiro ordenados por {campo_sp or 'orden del archivo'}"
            return np.array([(p[2], p[3]) for p in puntos]), nota
        raise ValueError("La línea no tiene geometría utilizable (líneas o puntos de tiro).")

    def _indice_nombres(self):
        """Índice (clave, campo, valor) de todos los nombres de línea del catálogo.

        La clave es la secuencia de bloques delimitada por '|' (p. ej. '|CV|1979|8|'). Solo
        se indexan valores con al menos una letra y un número, que es la forma de un nombre
        de línea; se excluyen cuencas, empresas o años sueltos. Se calcula una sola vez.
        """
        if getattr(self, "_indice", None) is None:
            vistos, self._indice = set(), []
            for c in self.campos_texto:
                for v in self._valores(c):
                    tk_ = tokens_nombre(v, ocr=True)
                    if (len(tk_) < 2 or not any(t.isdigit() for t in tk_)
                            or not any(t.isalpha() for t in tk_)):
                        continue
                    if (c, v) not in vistos:
                        vistos.add((c, v))
                        self._indice.append(("|" + "|".join(tk_) + "|", c, v))
        return self._indice

    def identificar(self, textos):
        """Identifica la línea a partir de los textos leídos en la imagen.

        Un nombre del catálogo se acepta solo si TODOS sus bloques aparecen consecutivos en
        un mismo texto ('|CV|1979|8|' dentro de '|LINEA|CV|1979|8|'). Los delimitadores
        impiden coincidencias parciales: 'CV-1979-08' no se confunde con 'CV-1979-080'.
        Si un nombre está contenido en otro también hallado (el programa 'CV-1979' dentro
        de la línea 'CV-1979-08'), se conserva solo el más específico.
        Devuelve [(campo, valor, texto_origen)] del más al menos específico.
        """
        cands = []
        for t in textos:
            for variante in (t, t.replace(" ", "")):        # también sin espacios: 'CVI 979-08'
                tk_ = tokens_nombre(variante, ocr=True)
                if tk_:
                    cands.append(("|" + "|".join(tk_) + "|", t))
        todo = "#".join(c for c, _ in cands)                # todos los textos en una cadena
        hallados = {}
        for clave, campo, valor in self._indice_nombres():
            if clave in todo:                               # el nombre aparece completo
                origen = next(t for c, t in cands if clave in c)   # texto donde se leyó
                hallados.setdefault(valor, (campo, clave, origen))
        claves = [c for _, c, _ in hallados.values()]
        res = [(campo, valor, origen, clave) for valor, (campo, clave, origen) in hallados.items()
               if not any(clave != c2 and clave in c2 for c2 in claves)]   # solo los más específicos
        res.sort(key=lambda r: (-len(r[3]), r[1]))          # primero las claves más largas
        return [(c, v, o) for c, v, o, _ in res]

    def atributos(self, campo, valor, max_campos=8):
        """Atributos del primer registro de la línea (para que el usuario la confirme)."""
        i = next(i for i, v in enumerate(self.datos[campo]) if str(v) == valor)
        n = sum(1 for v in self.datos[campo] if str(v) == valor)        # registros con ese nombre
        attrs = {c: self.datos[c][i] for c in self.campos
                 if self.datos[c][i] is not None and str(self.datos[c][i]).strip()
                 and str(self.datos[c][i]) != "nan"}
        return dict(list(attrs.items())[:max_campos]), n


# =========================================================================================
# 7. LECTURA DEL TEXTO DE LA IMAGEN (OCR integrado de Windows 10/11)
# =========================================================================================
def ocr_imagen(ruta):
    """Lee el texto impreso en la imagen con el motor OCR de Windows (Windows.Media.Ocr).

    Devuelve [(texto_de_la_línea, [(palabra, x, y, ancho, alto), ...]), ...] con posiciones
    en píxeles de la imagen original, o None si el OCR no está disponible en el equipo.
    La imagen se lee a varias escalas (x2, x3, x4) porque los rótulos de las secciones
    suelen ser pequeños; leer varias veces aumenta la probabilidad de un texto correcto.
    """
    try:   # las librerías winrt solo existen en Windows; en otro sistema se omite el OCR
        from winrt.windows.media.ocr import OcrEngine
        from winrt.windows.graphics.imaging import SoftwareBitmap, BitmapPixelFormat
        from winrt.windows.storage.streams import DataWriter
        from winrt.windows.globalization import Language
    except Exception:
        return None
    motor = OcrEngine.try_create_from_user_profile_languages()   # idioma de Windows del usuario
    if motor is None:                                            # si no, español o inglés
        for tag in ("es-ES", "es-MX", "en-US"):
            try:
                motor = OcrEngine.try_create_from_language(Language(tag))
            except Exception:
                motor = None
            if motor is not None:
                break
    if motor is None:
        return None
    im = Image.open(ruta).convert("RGBA")                 # el motor espera píxeles RGBA de 8 bits
    lado = max(im.size)
    limite = min(int(OcrEngine.max_image_dimension), 9000)        # tamaño máximo admitido
    escalas = sorted({round(min(s, limite / lado), 3) for s in (2, 3, 4)})   # escalas válidas
    salida = []
    for s in escalas:
        im2 = im.resize((max(1, int(im.width * s)), max(1, int(im.height * s))), Image.LANCZOS)
        w = DataWriter()                                  # se copian los píxeles a un búfer
        w.write_bytes(im2.tobytes())
        bmp = SoftwareBitmap.create_copy_from_buffer(w.detach_buffer(), BitmapPixelFormat.RGBA8,
                                                     im2.width, im2.height)
        res = asyncio.run(motor.recognize_async(bmp))     # reconocimiento (llamada asíncrona)
        for linea in res.lines:
            # posiciones de cada palabra devueltas a la escala de la imagen original
            palabras = [(p.text, p.bounding_rect.x / s, p.bounding_rect.y / s,
                         p.bounding_rect.width / s, p.bounding_rect.height / s) for p in linea.words]
            salida.append((linea.text, palabras))
    return salida


def cardinales_imagen(lineas_ocr, ancho):
    """Busca rótulos de orientación (SW, NE, O, E...) en el 25 % izquierdo y derecho de la imagen.

    Devuelve (rótulo_izquierdo, rótulo_derecho); si hay varios, el más alto de cada lado.
    """
    izq, der = [], []
    for _, palabras in lineas_ocr or []:
        for texto, x, y, w, _ in palabras:
            t = re.sub(r"[^A-Z]", "", _sin_tildes(texto).upper())   # solo letras
            if not RE_CARDINAL.match(t):                              # ¿es un rumbo válido?
                continue
            cx = x + w / 2                                            # centro horizontal
            if cx < ancho * 0.25:
                izq.append((y, t))
            elif cx > ancho * 0.75:
                der.append((y, t))
    return (min(izq)[1] if izq else None), (min(der)[1] if der else None)


# =========================================================================================
# 8. DESCARGA DEL CATÁLOGO DE SÍSMICA 2D DESDE EL SERVICIO ArcGIS REST DE LA ANH
# =========================================================================================
def _json_url(url, params, timeout=120, intentos=3):
    """Consulta un servicio ArcGIS REST y devuelve la respuesta JSON (con reintentos)."""
    req = urllib.request.Request(url + "?" + urllib.parse.urlencode(params),
                                 headers={"User-Agent": f"Mozilla/5.0 ({APP} {VERSION})"})
    ultimo = None
    for _ in range(intentos):
        try:
            with urllib.request.urlopen(req, timeout=timeout) as r:
                d = json.loads(r.read().decode("utf-8"))
            if "error" in d:                               # el servidor puede responder un error
                raise RuntimeError(d["error"].get("message", str(d["error"])))
            return d
        except (urllib.error.URLError, TimeoutError, OSError, ValueError) as e:
            ultimo = e                                     # se guarda y se reintenta
    raise RuntimeError(f"No hubo respuesta del servicio de la ANH ({ultimo}).")


def descargar_catalogo_anh(url, destino, avance=lambda hechos, total: None):
    """Descarga la capa Sísmica 2D completa a un GeoPackage local.

    La geometría se solicita en EPSG:3116 (sistema nativo de la capa) y sin simplificar, por
    lo que la copia local es idéntica a la publicada. El servicio entrega como máximo
    maxRecordCount registros por consulta, así que se pagina con resultOffset.
    """
    meta = _json_url(url, {"f": "json"})                           # descripción de la capa
    tipos = {f["name"]: f["type"] for f in meta["fields"] if f["type"] != "esriFieldTypeGeometry"}
    oid = meta.get("objectIdField") or next(                       # identificador para ordenar
        (n for n, t in tipos.items() if t == "esriFieldTypeOID"), None)
    total = _json_url(url + "/query", {"where": "1=1", "returnCountOnly": "true", "f": "json"})["count"]
    paso = min(int(meta.get("maxRecordCount") or 1000), 1000)      # registros por página
    feats, offset = [], 0
    while offset < total:                                          # paginación
        params = {"where": "1=1", "outFields": "*", "returnGeometry": "true", "outSR": "3116",
                  "resultOffset": offset, "resultRecordCount": paso, "f": "json"}
        if oid:
            params["orderByFields"] = oid                          # orden estable entre páginas
        lote = _json_url(url + "/query", params).get("features", [])
        if not lote:
            break
        feats.extend(lote)
        offset += len(lote)
        avance(offset, total)                                      # informe de progreso
    geoms, filas = [], []
    for ft in feats:                                               # formato Esri -> shapely
        caminos = [np.asarray(p, float)[:, :2] for p in (ft.get("geometry") or {}).get("paths", [])
                   if len(p) >= 2]
        if caminos:
            geoms.append(shapely.MultiLineString(caminos))
            filas.append(ft.get("attributes", {}))
    nombres, arreglos = [], []
    for n, t in tipos.items():                                     # conversión de tipos de campo
        vals = [f.get(n) for f in filas]
        if t in ("esriFieldTypeDouble", "esriFieldTypeSingle"):
            arr = np.array([np.nan if v is None else float(v) for v in vals], dtype=float)
        elif t in ("esriFieldTypeInteger", "esriFieldTypeSmallInteger", "esriFieldTypeOID"):
            arr = np.array([-1 if v is None else int(v) for v in vals], dtype=np.int64)
        elif t == "esriFieldTypeDate":                             # milisegundos desde 1970
            arr = np.array([None if v is None else (datetime.datetime(1970, 1, 1)
                            + datetime.timedelta(milliseconds=v)).strftime("%Y-%m-%d")
                            for v in vals], dtype=object)
        else:
            arr = np.array([None if v is None else str(v) for v in vals], dtype=object)
        nombres.append(re.sub(r"\W+", "_", n).strip("_"))          # nombres de campo válidos
        arreglos.append(arr)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    tmp = destino + ".tmp.gpkg"                                    # se escribe a un temporal y
    if os.path.exists(tmp):                                        # luego se reemplaza: una
        os.remove(tmp)                                             # descarga fallida no daña
    pyogrio.raw.write(tmp, np.array(shapely.to_wkb(geoms), dtype=object), arreglos, nombres,
                      layer="sismica_2d_anh", driver="GPKG", geometry_type="MultiLineString",
                      crs="EPSG:3116")                             # el catálogo anterior
    os.replace(tmp, destino)
    return len(geoms)


# =========================================================================================
# 9. CÁLCULO DE LOS PSEUDOPOZOS
# =========================================================================================
def calcular(cfg, linea):
    """Calcula posición y espesores de todos los pseudopozos.

    cfg   : diccionario con la calibración, los horizontes digitalizados y los parámetros.
    linea : LineaGeo ya orientada (su primer vértice es el extremo izquierdo de la sección).
    Devuelve (filas, avisos): una fila por pseudopozo y los avisos para el usuario.
    """
    # --- Calibración de la imagen: dos transformaciones lineales independientes ---
    xL, xR = cfg["xcal"]            # píxeles de los extremos izquierdo y derecho de la sección
    y1, y2 = cfg["ycal"]            # píxeles de las dos referencias verticales
    v1, v2 = cfg["ref"]             # valores de esas referencias (s TWT o km)
    L = cfg["long_km"]              # longitud de la línea en km
    esc_y = (v2 - v1) / (y2 - y1)                   # escala vertical (unidades por píxel)
    a_val = lambda py: v1 + (py - y1) * esc_y       # píxel vertical -> TWT o profundidad
    a_px = lambda km: xL + km / L * (xR - xL)       # km a lo largo de la línea -> píxel horizontal
    tol = TOL_PX * abs(esc_y)                       # tolerancia de espesor nulo, en unidades
    hor = cfg["horizontes"]
    twt = cfg["modo"] == "TWT"
    capas = cfg["capas"] if twt else []             # en profundidad no se separan capas
    filas, avisos = [], []
    for n, km in enumerate(cfg["pozos_km"], start=1):
        px = a_px(km)                                       # abscisa del pseudopozo en la imagen
        ys = evaluar(hor.get("sup", []), px)                # superficie (si se digitalizó)
        sup = a_val(ys) if np.isfinite(ys) else 0.0         # sin superficie: TWT = 0 o prof. = 0
        yb = evaluar(hor.get("bas", []), px)                # límite con el basamento
        if not np.isfinite(yb):                             # sin basamento no hay espesor total
            avisos.append(f"Pseudopozo {n} (km {km:.2f}): no hay límite de basamento digitalizado "
                          "en esa posición; se omitió.")
            continue
        bas = a_val(yb)
        if bas < sup:                                       # inconsistencia de digitalización
            avisos.append(f"Pseudopozo {n}: el basamento queda por encima de la superficie; espesor 0.")
            bas = sup
        # --- Posición sobre la línea BIP ---
        # Con "ajustar a longitud BIP" la sección completa (km 0 a L) se hace corresponder con
        # la línea BIP completa: la posición es la fracción km / L, independiente de la escala
        # gráfica de la imagen. En otro caso se usa el km BIP del extremo izquierdo como origen.
        frac = km / L if cfg["ajustar_bip"] else (cfg["offset_km"] + km) / linea.longitud_km
        if frac < -1e-9 or frac > 1 + 1e-9:
            avisos.append(f"Pseudopozo {n}: queda fuera de la navegación BIP "
                          f"({linea.longitud_km:.2f} km); se ubicó en el extremo.")
        p = linea.punto(frac)                               # coordenadas sobre la línea oficial
        f = dict(punto=cfg["punto_inicial"] + n - 1, pozo=n, linea=cfg["linea"],
                 lon=p["lon"], lat=p["lat"], este=p["este"], norte=p["norte"],
                 x_bip=p["x"], y_bip=p["y"], epsg_bip=linea.epsg, desv_m=p["desv_m"],
                 dist_km=km, dist_bip_km=min(max(frac, 0), 1) * linea.longitud_km, modo=cfg["modo"],
                 twt={}, vel={}, esp={}, twt_bas=None, esp_total=None)
        if twt:
            # --- Espesores por capa: se recorre la columna de techo a base ---
            tope = sup                                      # techo de la primera capa
            for k in ORDEN_CAPAS:
                f["vel"][k] = cfg["vel"][k]
                if k not in capas:                          # capa ausente en toda la línea
                    f["twt"][k], f["esp"][k] = None, 0.0
                    continue
                if k == capas[-1]:
                    base = bas          # la capa más profunda presente llega hasta el basamento
                else:
                    yh = evaluar(hor.get(k, []), px)        # base digitalizada de la capa
                    if not np.isfinite(yh):
                        base = tope     # capa ausente en esta posición (acuñada): espesor 0
                    else:
                        base = min(max(a_val(yh), tope), bas)   # entre su techo y el basamento
                        if bas - base <= tol:               # prácticamente sobre el basamento
                            base = bas
                        if base - tope <= tol:              # espesor menor que la tolerancia
                            base = tope
                dt = base - tope                            # intervalo de tiempo doble (s)
                f["twt"][k] = base if dt > 0 else None      # TWT de la base (vacío si no existe)
                f["esp"][k] = cfg["vel"][k] * dt / 2.0      # e = V * ΔTWT / 2
                tope = base                                 # la base es el techo de la siguiente
            f["twt_bas"] = bas
            f["esp_total"] = sum(f["esp"].values())         # espesor total = suma de capas
        else:
            f["esp_total"] = (bas - sup) * 1000.0           # profundidad en km -> espesor en m
        filas.append(f)
    return filas, avisos


# =========================================================================================
# 10. EXPORTACIÓN A EXCEL Y CSV
# =========================================================================================
# Hoja "Pseudopozos": mismo formato de la base de datos de puntos de control del proyecto.
ENC_PRINCIPAL = ["Punto de control", "Pseudopozo", "Línea", "Longitud", "Latitud",
                 "TWT Base Q", "TWT Base N", "TWT Base P",
                 "Velocidad Q (m/s)", "Velocidad N (m/s)", "Velocidad P (m/s)",
                 "Espesor Q (metros)", "Espesor N (metros)", "Espesor P (metros)",
                 "Espesor Total (metros)", "Distancia (km)", "TWT Basamento (s)",
                 "Tipo de interpretación"]
# Hoja "GIS": nombres sin tildes ni espacios, aptos como campos en QGIS y ArcGIS.
ENC_GIS = ["Punto", "Pseudopozo", "Linea", "Longitud", "Latitud", "Este_9377", "Norte_9377",
           "X_BIP", "Y_BIP", "EPSG_BIP", "Desv_BIP_m", "Dist_km", "Dist_BIP_km", "Tipo",
           "TWT_BQ", "TWT_BN", "TWT_BP", "TWT_Bas",
           "Vel_Q", "Vel_N", "Vel_P", "Esp_Q", "Esp_N", "Esp_P", "Esp_Total"]
# Hoja "Metadatos": trazabilidad de cada línea procesada.
ENC_META = ["Fecha", "Línea", "Tipo", "Identificación", "Archivo BIP", "Capa BIP", "Campo nombre",
            "EPSG BIP", "Imagen", "Longitud imagen (km)", "Longitud BIP (km)", "Ubicación",
            "Extremo izquierdo", "Desviación máx. respecto a línea BIP (m)",
            "Capas presentes", "Vel. Q (m/s)", "Vel. N (m/s)", "Vel. P (m/s)",
            "Pseudopozos", "Puntos de control"]
RELLENO = {"Q": "BFBFBF", "N": "FFFF66", "P": "F7A35C"}     # colores por capa (gris, amarillo, naranja)
COL_CAPA = {6: "Q", 7: "N", 8: "P", 9: "Q", 10: "N", 11: "P", 12: "Q", 13: "N", 14: "P"}  # columna -> capa
BORDE = Border(*(Side(style="thin", color="808080"),) * 4)  # borde fino en las cuatro caras


def _r(v, nd=4):
    """Redondea a nd decimales conservando los valores vacíos."""
    return None if v is None else round(float(v), nd)


def fila_principal(f):
    """Fila de la hoja Pseudopozos. Las secciones en profundidad llevan 'NA' por capa."""
    base = [f["punto"], f["pozo"], f["linea"], round(f["lon"], 8), round(f["lat"], 8)]  # ~1 mm
    if f["modo"] == "TWT":
        capas = ([_r(f["twt"][k]) if f["twt"][k] is not None else "" for k in ORDEN_CAPAS]
                 + [f["vel"][k] for k in ORDEN_CAPAS]
                 + [_r(f["esp"][k]) for k in ORDEN_CAPAS])
        twt_bas = _r(f["twt_bas"])
    else:
        capas = ["NA"] * 9          # no aplica conversión TWT -> espesor
        twt_bas = "NA"
    tipo = "TWT" if f["modo"] == "TWT" else "Profundidad (km)"
    return base + capas + [_r(f["esp_total"]), _r(f["dist_km"]), twt_bas, tipo]


def fila_gis(f):
    """Fila de la hoja GIS: valores numéricos (vacíos en lugar de 'NA' para no romper los tipos)."""
    twt = f["modo"] == "TWT"
    g = lambda d, k: (_r(d.get(k)) if twt else None)
    return [f["punto"], f["pozo"], f["linea"], round(f["lon"], 8), round(f["lat"], 8),
            round(f["este"], 3), round(f["norte"], 3), round(f["x_bip"], 3), round(f["y_bip"], 3),
            f["epsg_bip"], round(f["desv_m"], 3), _r(f["dist_km"]), _r(f["dist_bip_km"]),
            "TWT" if twt else "PROF",
            g(f["twt"], "Q"), g(f["twt"], "N"), g(f["twt"], "P"), _r(f["twt_bas"]) if twt else None,
            g(f["vel"], "Q"), g(f["vel"], "N"), g(f["vel"], "P"),
            g(f["esp"], "Q"), g(f["esp"], "N"), g(f["esp"], "P"), _r(f["esp_total"])]


def _encabezado(ws, enc, rellenos=None):
    """Escribe y da formato a la fila de encabezados de una hoja."""
    ws.append(enc)
    for j in range(1, len(enc) + 1):
        c = ws.cell(row=1, column=j)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE
        color = (rellenos or {}).get(j)
        if color:
            c.fill = PatternFill("solid", fgColor=color)
        ws.column_dimensions[get_column_letter(j)].width = max(11, min(24, len(enc[j - 1]) + 3))
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "A2"          # el encabezado queda fijo al desplazarse


def ultimo_punto(ruta):
    """Último número de punto de control de un Excel existente (para continuar la numeración)."""
    wb = load_workbook(ruta, read_only=True)
    if "Pseudopozos" not in wb.sheetnames:
        return None
    nums = [r[0] for r in wb["Pseudopozos"].iter_rows(min_row=2, max_col=1, values_only=True)
            if isinstance(r[0], (int, float))]
    wb.close()
    return int(max(nums)) if nums else 0


def exportar_excel(filas, meta, ruta, agregar=False):
    """Crea el Excel (o agrega filas a uno existente) y regenera el CSV de la hoja GIS."""
    if agregar and os.path.exists(ruta):
        wb = load_workbook(ruta)
        ws, wg, wm = wb["Pseudopozos"], wb["GIS"], wb["Metadatos"]
        # Se verifica que las columnas sean las de esta versión antes de agregar filas.
        if [c.value for c in wg[1]] != ENC_GIS or [c.value for c in wm[1]] != ENC_META:
            raise ValueError("El Excel fue creado con otra versión del programa (columnas distintas). "
                             "Guarde esta línea en un archivo nuevo.")
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Pseudopozos"
        wg = wb.create_sheet("GIS")
        wm = wb.create_sheet("Metadatos")
        _encabezado(ws, ENC_PRINCIPAL, {j: RELLENO[k] for j, k in COL_CAPA.items()})
        _encabezado(wg, ENC_GIS)
        _encabezado(wm, ENC_META)
    for f in filas:
        ws.append(fila_principal(f))
        r = ws.max_row
        for j in range(1, len(ENC_PRINCIPAL) + 1):          # formato de la fila recién escrita
            c = ws.cell(row=r, column=j)
            c.border = BORDE
            c.alignment = Alignment(horizontal="center")
            if j in COL_CAPA:
                c.fill = PatternFill("solid", fgColor=RELLENO[COL_CAPA[j]])
            elif 2 <= j <= 5:
                c.fill = PatternFill("solid", fgColor="E7E6E6")
        wg.append(fila_gis(f))
    wm.append([meta.get(k) for k in ENC_META])              # una fila de metadatos por línea
    wb.save(ruta)
    # CSV con toda la hoja GIS: UTF-8 con BOM (Excel reconoce las tildes), coma como separador
    # y punto decimal, que es lo que esperan QGIS y ArcGIS.
    ruta_csv = os.path.splitext(ruta)[0] + "_GIS.csv"
    with open(ruta_csv, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        for fila in wg.iter_rows(values_only=True):
            w.writerow(["" if v is None else v for v in fila])
    return ruta_csv


# =========================================================================================
# 11. INTERFAZ GRÁFICA
# =========================================================================================
AYUDA = """FLUJO DE TRABAJO

0. Catálogo BIP (una sola vez)
   "Descargar catálogo ANH" baja la capa oficial Sísmica 2D del geovisor de la ANH
   (geometría completa en MAGNA-SIRGAS Bogotá, EPSG:3116) y la guarda en el equipo.
   También puede usar la navegación que descargó del BIP (shapefile, GeoPackage, KML o
   carpeta .gdb). El catálogo queda recordado y se carga solo al abrir el programa.

1. Imagen de la sección -> identificación automática de la línea
   Al cargar la imagen, el programa lee su texto (OCR de Windows) y el nombre del archivo,
   y los cruza con el catálogo. Solo acepta nombres que EXISTEN en el BIP y cuyos bloques
   (p. ej. CV | 1979 | 08) aparecen completos y seguidos en el texto; corrige confusiones
   típicas del OCR (I/1, O/0). Se muestran los atributos de la línea para confirmarla.
   Si el texto no es legible, escriba el nombre en "Búsqueda manual".
   La orientación se toma de los rótulos de la sección (SW, NE...) si el OCR los lee; si
   no, el programa pregunta cuál extremo de la línea BIP está a la izquierda.

2. Escala horizontal
   Marque con dos clics el extremo izquierdo y el derecho de la sección. Con "Ajustar a
   longitud BIP" (recomendado) la sección completa corresponde a la línea BIP completa:
   cada pseudopozo se ubica por su fracción de la longitud de la línea, independiente
   de la escala gráfica de la imagen. Desactívelo solo si la imagen muestra un tramo de
   la línea, e indique el km BIP de su extremo izquierdo.

EXACTITUD GEOESPACIAL
   Los puntos se calculan sobre los vértices de la línea BIP en su sistema original,
   sin reproyectarla antes, así que quedan exactamente sobre la geometría oficial. Cada
   punto se verifica contra la línea (columna Desv_BIP_m, que debe ser 0.000) y se
   exportan también sus coordenadas originales (X_BIP, Y_BIP). La única incertidumbre
   restante es a lo largo de la línea y depende de marcar bien los extremos de la sección.

3. Escala vertical
   Escriba los valores de dos referencias (p. ej. 0 y 1 s TWT, o 0 y 1 km) y márquelas con dos clics.

4. Capas (solo TWT)
   Indique cuántas capas tiene la línea y cuáles son:
   Capa 1 = Cuaternario, Capa 2 = Neógeno, Capa 3 = Paleógeno.
   La base de la capa más profunda presente es el límite con basamento, por eso no se
   digitaliza aparte. Si una capa se acuña, digitalice su base solo donde exista; donde
   no haya horizonte la capa tiene espesor 0. Los vértices cercanos a otro horizonte se
   pegan a él (imán), de modo que los acuñamientos quedan exactos.

5. Digitalización
   Clic izquierdo agrega vértices, clic derecho deshace, doble clic o Enter termina.
   Rueda del ratón: acercar/alejar. Use la mano de la barra para desplazarse
   (desactívela para volver a digitalizar). Sin superficie digitalizada, el tope es TWT = 0
   (o profundidad 0).

6. Pseudopozos
   Defina la cantidad o el espaciado (km) y el rango. Pulse Previsualizar.

7. Calcular y exportar
   Genera Excel (hojas Pseudopozos, GIS y Metadatos), un CSV para SIG, una imagen de
   control y el proyecto (.json) para retomar la digitalización. Si elige un Excel
   existente puede AGREGAR la línea a la base; la numeración de puntos continúa.

Coordenadas: Longitud/Latitud en grados decimales WGS84 (EPSG:4326) y Este/Norte en
MAGNA-SIRGAS Origen Nacional (EPSG:9377). En QGIS: Capa > Añadir capa de texto
delimitado (CSV) con X = Longitud, Y = Latitud, EPSG:4326. En ArcGIS Pro: XY Table To
Point sobre la hoja GIS o el CSV.
"""


class App(tk.Tk):
    """Ventana principal: panel de pasos a la izquierda y la sección sísmica a la derecha."""

    def __init__(self):
        super().__init__()
        self.title(f"{APP} {VERSION}")
        self.geometry("1450x880")
        try:
            self.state("zoomed")                    # ventana maximizada en Windows
        except tk.TclError:
            pass
        self.report_callback_exception = self._error   # errores no previstos -> mensaje legible

        # --- Estado del trabajo en curso ---
        self.bip = None              # catálogo BIP cargado (BaseBIP)
        self.linea_base = None       # geometría de la línea identificada (LineaGeo)
        self.linea_valor = None      # nombre de la línea tal como figura en el BIP
        self.linea_campo = None      # campo del catálogo que contiene ese nombre
        self.img = None              # imagen de la sección como matriz de píxeles
        self.img_ruta = None         # ruta del archivo de imagen
        self.xcal, self.ycal = [], []   # píxeles de calibración horizontal y vertical
        self.hor = {"sup": [], "Q": [], "N": [], "P": [], "bas": []}   # horizontes digitalizados
        self.activo = None           # qué se está marcando con el ratón (calibración u horizonte)
        self.pozos_km = []           # posiciones de los pseudopozos (km desde el extremo izquierdo)
        self.capas = ["Q", "N", "P"] # capas presentes en la línea
        self.capas_definidas = False # el usuario confirmó las capas
        self._artistas = []          # elementos dibujados sobre la imagen (para redibujar)
        self._ocr = None             # resultado del OCR de la imagen actual
        self.identificacion = ""     # cómo se identificó la línea (trazabilidad)
        self._descarga = None        # hilo de descarga del catálogo ANH

        # --- Variables enlazadas a los controles de la interfaz ---
        self.v_archivo = tk.StringVar(value="Sin catálogo BIP cargado.")
        self.v_campo = tk.StringVar()
        self.v_nombre = tk.StringVar()
        self.v_info_linea = tk.StringVar(value="Sin línea seleccionada.")
        self.v_modo = tk.StringVar(value="TWT")                 # TWT o profundidad (PROF)
        self.v_long = tk.StringVar()                            # longitud de la línea (km)
        self.v_ajustar = tk.BooleanVar(value=True)              # sección completa = línea completa
        self.v_offset = tk.StringVar(value="0")                 # km BIP del extremo izquierdo
        self.v_extremo = tk.StringVar(value="Primer vértice BIP")
        self.v_ref1 = tk.StringVar(value="0")                   # valor de la referencia vertical 1
        self.v_ref2 = tk.StringVar(value="1")                   # valor de la referencia vertical 2
        self.v_unid = tk.StringVar(value="s (TWT)")
        self.v_vel = {k: tk.StringVar(value=str(CAPAS[k][2])) for k in ORDEN_CAPAS}
        self.v_capas_txt = tk.StringVar(value="Capas sin definir.")
        self.v_tipo_pozos = tk.StringVar(value="cantidad")      # cantidad o espaciado
        self.v_val_pozos = tk.StringVar(value="34")
        self.v_km_ini = tk.StringVar(value="0")
        self.v_km_fin = tk.StringVar()
        self.v_punto_ini = tk.StringVar(value="1")              # primer número de punto de control
        self.v_snap = tk.BooleanVar(value=True)                 # imán entre horizontes
        self.v_msg = tk.StringVar(value="Cargue la imagen de la sección sísmica (paso 1).")
        self._construir()
        self.after(200, self._catalogo_inicial)   # el catálogo se carga cuando la ventana ya existe

    # ------------------------------------------------------------------ construcción de la UI
    def _construir(self):
        """Crea el panel de pasos (con barra de desplazamiento) y el lienzo de la sección."""
        cont = ttk.Frame(self)
        cont.pack(side="left", fill="y")
        cv = tk.Canvas(cont, width=390, highlightthickness=0)    # contenedor desplazable
        sb = ttk.Scrollbar(cont, orient="vertical", command=cv.yview)
        panel = ttk.Frame(cv, padding=6)
        panel.bind("<Configure>", lambda e: cv.configure(scrollregion=cv.bbox("all")))
        cv.create_window((0, 0), window=panel, anchor="nw")
        cv.configure(yscrollcommand=sb.set)
        cv.pack(side="left", fill="y")
        sb.pack(side="left", fill="y")
        # La rueda del ratón desplaza el panel solo cuando el puntero está sobre él.
        rueda = lambda e: cv.yview_scroll(int(-e.delta / 120), "units")
        cont.bind("<Enter>", lambda e: cv.bind_all("<MouseWheel>", rueda))
        cont.bind("<Leave>", lambda e: cv.unbind_all("<MouseWheel>"))
        W = 360   # ancho de ajuste de los textos del panel

        # Paso 0. Catálogo BIP
        f = ttk.LabelFrame(panel, text="Catálogo BIP de sísmica 2D (una sola vez)", padding=6)
        f.pack(fill="x", pady=3)
        ttk.Label(f, textvariable=self.v_archivo, wraplength=W, foreground="#555").pack(anchor="w")
        fb = ttk.Frame(f)
        fb.pack(fill="x", pady=(3, 0))
        self.b_anh = ttk.Button(fb, text="Descargar catálogo ANH", command=self._descargar_anh)
        self.b_anh.pack(side="left")
        ttk.Button(fb, text="Archivo…", command=self._cargar_bip_archivo).pack(side="left", padx=3)
        ttk.Button(fb, text=".gdb…", command=self._cargar_bip_gdb).pack(side="left")

        # Paso 1. Imagen e identificación de la línea
        f = ttk.LabelFrame(panel, text="1. Imagen de la sección → línea BIP", padding=6)
        f.pack(fill="x", pady=3)
        fb = ttk.Frame(f)
        fb.pack(fill="x")
        ttk.Button(fb, text="Cargar imagen…", command=self._cargar_imagen).pack(side="left")
        ttk.Button(fb, text="Reidentificar", command=self._identificar).pack(side="left", padx=4)
        fm = ttk.Frame(f)
        fm.pack(fill="x", pady=2)
        ttk.Label(fm, text="Interpretación en:").pack(side="left")
        ttk.Radiobutton(fm, text="TWT (s)", variable=self.v_modo, value="TWT",
                        command=self._cambio_modo).pack(side="left")
        ttk.Radiobutton(fm, text="Profundidad (km)", variable=self.v_modo, value="PROF",
                        command=self._cambio_modo).pack(side="left")
        ttk.Label(f, textvariable=self.v_info_linea, wraplength=W, foreground="#0b5394").pack(anchor="w")
        fman = ttk.LabelFrame(f, text="Búsqueda manual (si el texto no es legible)", padding=4)
        fman.pack(fill="x", pady=(4, 0))
        fc = ttk.Frame(fman)
        fc.pack(fill="x")
        ttk.Label(fc, text="Campo:").pack(side="left")
        self.cb_campo = ttk.Combobox(fc, textvariable=self.v_campo, state="readonly", width=26)
        self.cb_campo.pack(side="left", padx=4)
        fn = ttk.Frame(fman)
        fn.pack(fill="x", pady=2)
        ttk.Label(fn, text="Nombre:").pack(side="left")
        en = ttk.Entry(fn, textvariable=self.v_nombre, width=20)
        en.pack(side="left", padx=4)
        en.bind("<Return>", lambda e: self._buscar_linea())
        ttk.Button(fn, text="Buscar", command=self._buscar_linea).pack(side="left")

        # Paso 2. Escala horizontal
        f = ttk.LabelFrame(panel, text="2. Escala horizontal", padding=6)
        f.pack(fill="x", pady=3)
        g = ttk.Frame(f)
        g.pack(fill="x")
        ttk.Label(g, text="Longitud de la línea (km):").grid(row=0, column=0, sticky="w")
        ttk.Entry(g, textvariable=self.v_long, width=10).grid(row=0, column=1, sticky="w")
        ttk.Label(g, text="Extremo izquierdo:").grid(row=1, column=0, sticky="w")
        ttk.Combobox(g, textvariable=self.v_extremo, values=EXTREMOS, state="readonly",
                     width=17).grid(row=1, column=1, sticky="w")
        ttk.Checkbutton(g, text="Ajustar a longitud BIP", variable=self.v_ajustar).grid(
            row=2, column=0, columnspan=2, sticky="w")
        ttk.Label(g, text="km BIP del extremo izq.:").grid(row=3, column=0, sticky="w")
        ttk.Entry(g, textvariable=self.v_offset, width=10).grid(row=3, column=1, sticky="w")
        ttk.Button(f, text="Marcar extremos izq./der. (2 clics)",
                   command=lambda: self._activar("xcal")).pack(anchor="w", pady=(4, 0))

        # Paso 3. Escala vertical
        f = ttk.LabelFrame(panel, text="3. Escala vertical", padding=6)
        f.pack(fill="x", pady=3)
        g = ttk.Frame(f)
        g.pack(fill="x")
        ttk.Label(g, text="Referencia 1:").grid(row=0, column=0, sticky="w")
        ttk.Entry(g, textvariable=self.v_ref1, width=8).grid(row=0, column=1)
        ttk.Label(g, textvariable=self.v_unid).grid(row=0, column=2, sticky="w", padx=3)
        ttk.Label(g, text="Referencia 2:").grid(row=1, column=0, sticky="w")
        ttk.Entry(g, textvariable=self.v_ref2, width=8).grid(row=1, column=1)
        ttk.Label(g, textvariable=self.v_unid).grid(row=1, column=2, sticky="w", padx=3)
        ttk.Button(f, text="Marcar referencias 1 y 2 (2 clics)",
                   command=lambda: self._activar("ycal")).pack(anchor="w", pady=(4, 0))

        # Paso 4. Capas y velocidades interválicas
        self.f_capas = ttk.LabelFrame(panel, text="4. Capas y velocidades interválicas (TWT)", padding=6)
        self.f_capas.pack(fill="x", pady=3)
        ttk.Button(self.f_capas, text="Definir capas de la línea…",
                   command=self._dialogo_capas).pack(anchor="w")
        ttk.Label(self.f_capas, textvariable=self.v_capas_txt, wraplength=W).pack(anchor="w")
        g = ttk.Frame(self.f_capas)
        g.pack(fill="x", pady=2)
        for i, k in enumerate(ORDEN_CAPAS):                 # una casilla de velocidad por capa
            n, nombre, _, _ = CAPAS[k]
            ttk.Label(g, text=f"Capa {n} – {nombre} ({k}) m/s:").grid(row=i, column=0, sticky="w")
            ttk.Entry(g, textvariable=self.v_vel[k], width=9).grid(row=i, column=1)

        # Paso 5. Digitalización de horizontes (los botones dependen de las capas elegidas)
        f = ttk.LabelFrame(panel, text="5. Digitalización de horizontes", padding=6)
        f.pack(fill="x", pady=3)
        self.f_hor = ttk.Frame(f)
        self.f_hor.pack(fill="x")
        g = ttk.Frame(f)
        g.pack(fill="x", pady=(4, 0))
        ttk.Button(g, text="Terminar", command=self._terminar).pack(side="left")
        ttk.Button(g, text="Deshacer punto", command=self._deshacer).pack(side="left", padx=3)
        ttk.Button(g, text="Borrar horizonte", command=self._borrar_horizonte).pack(side="left")
        ttk.Checkbutton(f, text="Imán: pegar vértices a otros horizontes",
                        variable=self.v_snap).pack(anchor="w")

        # Paso 6. Pseudopozos
        f = ttk.LabelFrame(panel, text="6. Pseudopozos", padding=6)
        f.pack(fill="x", pady=3)
        g = ttk.Frame(f)
        g.pack(fill="x")
        ttk.Radiobutton(g, text="Cantidad", variable=self.v_tipo_pozos, value="cantidad").grid(
            row=0, column=0, sticky="w")
        ttk.Radiobutton(g, text="Espaciado (km)", variable=self.v_tipo_pozos, value="espaciado").grid(
            row=0, column=1, sticky="w")
        ttk.Entry(g, textvariable=self.v_val_pozos, width=8).grid(row=0, column=2)
        ttk.Label(g, text="Desde km:").grid(row=1, column=0, sticky="w")
        ttk.Entry(g, textvariable=self.v_km_ini, width=8).grid(row=1, column=1, sticky="w")
        ttk.Label(g, text="Hasta km:").grid(row=2, column=0, sticky="w")
        ttk.Entry(g, textvariable=self.v_km_fin, width=8).grid(row=2, column=1, sticky="w")
        ttk.Label(g, text="Primer punto de control:").grid(row=3, column=0, columnspan=2, sticky="w")
        ttk.Entry(g, textvariable=self.v_punto_ini, width=8).grid(row=3, column=2)
        ttk.Button(f, text="Previsualizar pseudopozos", command=self._generar_pozos).pack(
            anchor="w", pady=(4, 0))

        # Paso 7. Resultados
        f = ttk.LabelFrame(panel, text="7. Resultados", padding=6)
        f.pack(fill="x", pady=3)
        ttk.Button(f, text="Calcular y exportar Excel…", command=self._exportar).pack(fill="x")
        g = ttk.Frame(f)
        g.pack(fill="x", pady=3)
        ttk.Button(g, text="Guardar proyecto…", command=self._guardar_proyecto).pack(side="left")
        ttk.Button(g, text="Abrir proyecto…", command=self._abrir_proyecto).pack(side="left", padx=3)
        ttk.Button(g, text="Ayuda", command=self._ayuda).pack(side="left")

        # Lienzo de la sección: barra de instrucciones, figura de matplotlib y barra de zoom.
        der = ttk.Frame(self)
        der.pack(side="left", fill="both", expand=True)
        ttk.Label(der, textvariable=self.v_msg, background="#fff8dc", padding=6,
                  wraplength=1000, font=("Segoe UI", 10)).pack(fill="x")
        self.fig = Figure(figsize=(10, 6))
        self.ax = self.fig.add_axes([0, 0, 1, 1])     # ejes que ocupan toda la figura
        self.ax.set_axis_off()                        # sin marcos: coordenadas en píxeles de imagen
        self.canvas = FigureCanvasTkAgg(self.fig, master=der)
        self.toolbar = NavigationToolbar2Tk(self.canvas, der)
        self.toolbar.update()
        self.canvas.get_tk_widget().pack(fill="both", expand=True)
        self.canvas.mpl_connect("button_press_event", self._clic)    # clics de digitalización
        self.canvas.mpl_connect("key_press_event", self._tecla)      # Enter, Esc, retroceso
        self.canvas.mpl_connect("scroll_event", self._rueda)         # zoom con la rueda
        self._rehacer_botones()

    def _error(self, exc, val, tb):
        """Muestra cualquier error inesperado en una ventana en lugar de cerrar el programa."""
        messagebox.showerror("Error", "".join(traceback.format_exception(exc, val, tb))[-2500:])

    def _ayuda(self):
        """Ventana con el texto de ayuda (flujo de trabajo)."""
        top = tk.Toplevel(self)
        top.title("Ayuda")
        t = tk.Text(top, wrap="word", width=95, height=40, font=("Segoe UI", 10))
        t.insert("1.0", AYUDA)
        t.configure(state="disabled")
        t.pack(fill="both", expand=True)

    def _elegir(self, titulo, texto, opciones):
        """Diálogo modal para elegir una opción de una lista; devuelve la elegida o None."""
        res = {"v": None}
        top = tk.Toplevel(self)
        top.title(titulo)
        top.transient(self)
        ttk.Label(top, text=texto, wraplength=420, padding=8).pack(anchor="w")
        lb = tk.Listbox(top, width=60, height=min(15, max(4, len(opciones))))
        for o in opciones:
            lb.insert("end", o)
        lb.pack(padx=8, fill="both", expand=True)
        if opciones:
            lb.selection_set(0)

        def ok(_=None):
            sel = lb.curselection()
            res["v"] = opciones[sel[0]] if sel else None
            top.destroy()
        lb.bind("<Double-Button-1>", ok)
        g = ttk.Frame(top, padding=8)
        g.pack()
        ttk.Button(g, text="Aceptar", command=ok).pack(side="left", padx=4)
        ttk.Button(g, text="Cancelar", command=top.destroy).pack(side="left")
        top.grab_set()               # bloquea la ventana principal hasta responder
        self.wait_window(top)
        return res["v"]

    # ------------------------------------------------------------------ catálogo BIP
    def _cargar_bip_archivo(self):
        """Carga la navegación del BIP desde un archivo vectorial."""
        ruta = filedialog.askopenfilename(
            title="Navegación sísmica 2D del BIP",
            filetypes=[("Datos vectoriales", "*.shp *.gpkg *.geojson *.json *.kml *.kmz *.gml"),
                       ("Todos", "*.*")])
        if ruta:
            self._cargar_bip(ruta)

    def _cargar_bip_gdb(self):
        """Carga la navegación del BIP desde una geodatabase de archivos (.gdb es una carpeta)."""
        ruta = filedialog.askdirectory(title="Seleccione la carpeta .gdb del BIP")
        if ruta:
            self._cargar_bip(ruta)

    def _catalogo_inicial(self):
        """Al iniciar, abre el catálogo usado la vez anterior (o la copia ANH descargada)."""
        c = leer_cfg().get("catalogo") or {}
        ruta = c.get("ruta")
        if not ruta and os.path.exists(RUTA_CATALOGO_ANH):
            ruta, c = RUTA_CATALOGO_ANH, {"capa": "sismica_2d_anh"}
        if ruta and os.path.exists(ruta):
            try:
                self._cargar_bip(ruta, c.get("capa"), c.get("campo"))
                return
            except Exception as e:
                self.v_archivo.set(f"No se pudo abrir el catálogo guardado ({e}).")
        self.v_msg.set("No hay catálogo BIP configurado. Pulse 'Descargar catálogo ANH' o cargue la "
                       "navegación sísmica 2D del BIP con 'Archivo…' / '.gdb…' (solo se hace una vez).")

    def _descargar_anh(self):
        """Descarga el catálogo ANH en un hilo aparte para no congelar la ventana."""
        if self._descarga is not None:          # ya hay una descarga en curso
            return
        url = leer_cfg().get("url_anh", URL_ANH)   # la URL puede cambiarse en config.json
        if not messagebox.askokcancel(APP, "Se descargará la capa Sísmica 2D del geovisor de la ANH:\n"
                                      f"{url}\n\nPuede tardar varios minutos. ¿Continuar?"):
            return
        estado = {"hechos": 0, "total": None, "fin": False, "error": None, "n": 0}

        def trabajo():                          # se ejecuta en segundo plano
            try:
                estado["n"] = descargar_catalogo_anh(
                    url, RUTA_CATALOGO_ANH, lambda h, t: estado.update(hechos=h, total=t))
            except Exception as e:
                estado["error"] = e
            estado["fin"] = True
        self._descarga = threading.Thread(target=trabajo, daemon=True)
        self._descarga.start()
        self.b_anh.configure(state="disabled")
        self._vigilar_descarga(estado)

    def _vigilar_descarga(self, estado):
        """Actualiza el progreso cada 0,4 s; Tkinter solo puede modificarse desde su propio hilo."""
        if not estado["fin"]:
            t = estado["total"]
            self.v_archivo.set(f"Descargando catálogo ANH… {estado['hechos']}"
                               + (f" de {t} registros" if t else ""))
            self.after(400, self._vigilar_descarga, estado)
            return
        self._descarga = None
        self.b_anh.configure(state="normal")
        if estado["error"]:
            self.v_archivo.set("La descarga del catálogo ANH falló.")
            messagebox.showerror(APP, f"{estado['error']}\n\nSi el geovisor de la ANH no responde, "
                                      "descargue la navegación sísmica 2D desde el BIP y cárguela con "
                                      "'Archivo…' o '.gdb…'.")
            return
        self._cargar_bip(RUTA_CATALOGO_ANH, "sismica_2d_anh")
        messagebox.showinfo(APP, f"Catálogo ANH descargado: {estado['n']} registros de sísmica 2D.\n"
                                 f"Queda guardado en:\n{RUTA_CATALOGO_ANH}")

    def _cargar_bip(self, ruta, capa=None, campo=None, auto_identificar=True):
        """Abre el catálogo, lo recuerda en la configuración y, si hay imagen, identifica la línea."""
        self.config(cursor="watch")
        self.update()
        try:
            if capa is None:                              # archivos con varias capas: elegir
                capas = BaseBIP.capas(ruta)
                if len(capas) > 1:
                    ops = [f"{n}   [{t}]" for n, t in capas]
                    sel = self._elegir("Capa", "El archivo tiene varias capas. Elija la de navegación "
                                       "sísmica 2D (líneas o puntos de tiro):", ops)
                    if sel is None:
                        return
                    capa = capas[ops.index(sel)][0]
                elif capas:
                    capa = capas[0][0]
            self.bip = BaseBIP(ruta, capa)
            self.bip._indice_nombres()                    # índice de nombres listo de antemano
        finally:
            self.config(cursor="")
        self.cb_campo["values"] = self.bip.campos_texto or self.bip.campos
        self.v_campo.set(campo if campo in self.bip.campos else self.bip.campo_sugerido())
        guardar_cfg(catalogo={"ruta": ruta, "capa": capa, "campo": self.v_campo.get()})
        self.v_archivo.set(f"{os.path.basename(ruta)}{' / ' + capa if capa else ''} — "
                           f"{len(self.bip.geoms)} registros — {self.bip.crs or 'sin CRS'}")
        if auto_identificar and self.img is not None and self.linea_base is None:
            self._identificar()
        elif self.linea_base is None:
            self.v_msg.set("Catálogo BIP listo. Paso 1: cargue la imagen de la sección.")

    def _crs_bip(self, xy):
        """Sistema de referencia del catálogo; si el archivo no lo declara se pregunta al usuario."""
        if self.bip.crs:
            return self.bip.crs
        if np.all(np.abs(xy[:, 0]) <= 180) and np.all(np.abs(xy[:, 1]) <= 90):
            return "EPSG:4326"                            # coordenadas que solo pueden ser grados
        epsg = simpledialog.askinteger(
            "Sistema de referencia", "El catálogo no declara sistema de coordenadas.\n"
            "Ingrese el código EPSG (p. ej. 3116 MAGNA Bogotá, 9377 Origen Nacional, "
            "21897 Bogotá 1975):", parent=self)
        return f"EPSG:{epsg}" if epsg else None

    def _fijar_linea(self, campo, valor, origen, orientar=True):
        """Carga la geometría de la línea elegida y la deja lista para el cálculo."""
        xy, nota = self.bip.geometria(campo, valor)
        crs = self._crs_bip(xy)
        if not crs:
            return False
        self.linea_base = LineaGeo(xy, crs)
        self.linea_valor, self.linea_campo, self.identificacion = valor, campo, origen
        self.v_campo.set(campo)
        self.v_nombre.set(valor)
        self.v_long.set(f"{self.linea_base.longitud_km:.3f}")   # longitud tomada del BIP
        self.v_km_fin.set(self.v_long.get())
        r0, r1 = self.linea_base.rumbo()
        self.v_info_linea.set(
            f"✔ {valor} — {self.linea_base.longitud_km:.3f} km ({nota}); va del {r0} al {r1}; "
            f"{self.linea_base.epsg}. Identificada por {origen}.")
        if orientar:
            self._orientar()
        self._dibujar()
        return True

    def _buscar_linea(self):
        """Búsqueda manual por nombre (alternativa cuando el texto de la imagen no es legible)."""
        if self.bip is None:
            messagebox.showwarning(APP, "Primero cargue el catálogo BIP.")
            return
        nombre = self.v_nombre.get().strip()
        if not nombre:
            return
        campo, exactos, sug = self.bip.buscar(nombre, self.v_campo.get())
        if exactos:
            valor = exactos[0] if len(exactos) == 1 else self._elegir(
                "Coincidencias", "Varias entradas coinciden:", exactos)
        elif sug:
            valor = self._elegir("Línea no encontrada", f"No hay coincidencia exacta para '{nombre}' "
                                 f"en el campo {campo}. ¿Es alguna de estas?", sug)
        else:
            messagebox.showwarning(APP, f"No se encontró '{nombre}' en el catálogo BIP. "
                                   "Revise el campo seleccionado o el nombre.")
            return
        if valor is not None:
            self._fijar_linea(campo, valor, f"búsqueda manual «{nombre}»")

    def _identificar(self):
        """Identifica la línea BIP a partir del texto de la imagen y del nombre del archivo."""
        if self.img_ruta is None:
            messagebox.showwarning(APP, "Primero cargue la imagen de la sección.")
            return
        if self.bip is None:
            messagebox.showwarning(APP, "Para identificar la línea se necesita el catálogo BIP: pulse "
                                   "'Descargar catálogo ANH' o cargue la navegación del BIP.")
            return
        if self._ocr is None:                           # el OCR se ejecuta una vez por imagen
            self.v_msg.set("Leyendo el texto de la imagen…")
            self.config(cursor="watch")
            self.update()
            try:
                leido = ocr_imagen(self.img_ruta)
            finally:
                self.config(cursor="")
            if leido is None:
                messagebox.showwarning(APP, "El OCR de Windows no está disponible en este equipo; "
                                       "solo se usará el nombre del archivo.")
            self._ocr = leido or []
        archivo = os.path.splitext(os.path.basename(self.img_ruta))[0]
        textos = [archivo] + [t for t, _ in self._ocr]  # nombre de archivo + textos leídos
        hallados = self.bip.identificar(textos)
        if not hallados:
            leidos = sorted({t for t in textos if any(ch.isdigit() for ch in t)})[:15]
            messagebox.showwarning(
                APP, "Ningún nombre de línea del catálogo BIP aparece en el texto de la imagen ni en "
                     "el nombre del archivo.\n\nTextos con números leídos:\n"
                     + ("\n".join(leidos) or "(ninguno)") + "\n\nUse la búsqueda manual.")
            self.v_msg.set("Línea no identificada automáticamente: use la búsqueda manual (paso 1).")
            return
        if len(hallados) > 1:                           # varias candidatas: decide el usuario
            ops = [f"{v}   (campo {c}; leído: «{o}»)" for c, v, o in hallados]
            sel = self._elegir("Varias líneas posibles", "El texto de la imagen coincide con varias "
                               "líneas del BIP. Elija la de esta sección:", ops)
            if sel is None:
                return
            campo, valor, origen = hallados[ops.index(sel)]
        else:                                           # una sola: se pide confirmación
            campo, valor, origen = hallados[0]
            attrs, n = self.bip.atributos(campo, valor)
            fuente = "nombre del archivo" if origen == archivo else "texto de la imagen"
            txt = (f"Línea identificada en el BIP:  {valor}\nLeída en: «{origen}»  ({fuente})\n\n"
                   + "\n".join(f"{k}: {v}" for k, v in attrs.items())
                   + (f"\n\n({n} registros con este nombre)" if n > 1 else "")
                   + "\n\n¿Es la línea de esta sección?")
            if not messagebox.askyesno("Línea identificada", txt):
                self.v_msg.set("Identificación rechazada: use la búsqueda manual (paso 1).")
                return
        fuente = "nombre del archivo" if origen == archivo else "texto de la imagen"
        self._fijar_linea(campo, valor, f"{fuente} «{origen}»")

    def _orientar(self):
        """Define qué extremo de la línea BIP está a la izquierda de la sección.

        Si el OCR leyó un rótulo de orientación a un lado de la imagen y su dirección forma un
        ángulo de menos de 60° (coseno >= 0,5) con uno de los extremos de la línea, se adopta
        automáticamente. En caso contrario se pregunta al usuario.
        """
        r0, r1 = self.linea_base.rumbo()
        izq, der = (cardinales_imagen(self._ocr, self.img.shape[1])
                    if self._ocr and self.img is not None else (None, None))
        for etiqueta, signo, lado in ((izq, 1, "izquierda"), (der, -1, "derecha")):
            if etiqueta:
                c = self.linea_base.coseno_primer_vertice(etiqueta) * signo
                if abs(c) >= 0.5:
                    primero = c > 0
                    self.v_extremo.set("Primer vértice BIP" if primero else "Último vértice BIP")
                    self.v_msg.set(f"Orientación leída en la imagen ({etiqueta} a la {lado}): extremo "
                                   f"izquierdo = {r0 if primero else r1}. Paso 2: marque los extremos "
                                   "de la sección.")
                    return
        res = {"v": None}
        top = tk.Toplevel(self)
        top.title("Orientación de la sección")
        top.transient(self)
        ttk.Label(top, padding=10, wraplength=440, text=(
            f"En el BIP la línea {self.linea_valor} va del {r0} (primer vértice) al {r1} "
            "(último vértice).\n\n¿Qué extremo está a la IZQUIERDA de la sección?\n"
            "Guíese por los rótulos de orientación de la imagen (p. ej. SW … NE).")).pack()
        g = ttk.Frame(top, padding=8)
        g.pack()
        for texto, val in ((f"{r0}  (primer vértice)", "Primer vértice BIP"),
                           (f"{r1}  (último vértice)", "Último vértice BIP")):
            ttk.Button(g, text=texto, command=lambda v=val: (res.update(v=v), top.destroy())).pack(
                side="left", padx=6)
        top.grab_set()
        self.wait_window(top)
        if res["v"]:
            self.v_extremo.set(res["v"])
        self.v_msg.set("Paso 2: pulse 'Marcar extremos' y haga clic en el extremo izquierdo y derecho "
                       "de la sección.")

    # ------------------------------------------------------------------ imagen
    def _cargar_imagen(self, ruta=None, identificar=True):
        """Carga la imagen de la sección y, si se pide, identifica la línea automáticamente."""
        ruta = ruta or filedialog.askopenfilename(
            title="Imagen de la sección sísmica",
            filetypes=[("Imágenes", "*.png *.jpg *.jpeg *.tif *.tiff *.bmp"), ("Todos", "*.*")])
        if not ruta:
            return
        self.img = np.asarray(Image.open(ruta).convert("RGB"))   # matriz alto x ancho x 3
        self.img_ruta = ruta
        self._ocr = None
        if identificar:   # imagen nueva: se descarta lo calibrado y digitalizado en la anterior
            self.linea_base = self.linea_valor = self.linea_campo = None
            self.v_info_linea.set("Sin línea identificada.")
            self.xcal, self.ycal, self.pozos_km = [], [], []
            self.hor = {k: [] for k in self.hor}
        self._artistas = []
        self.ax.clear()
        self.ax.set_axis_off()
        self.ax.imshow(self.img, interpolation="nearest")   # cada píxel = una unidad de los ejes
        self.ax.set_autoscale_on(False)                     # lo dibujado no altera el encuadre
        self.toolbar.update()
        self._dibujar()
        if identificar:
            if self.bip is None:
                self.v_msg.set("Imagen cargada. Falta el catálogo BIP para identificar la línea: pulse "
                               "'Descargar catálogo ANH' o cargue la navegación del BIP.")
            else:
                self._identificar()

    def _cambio_modo(self):
        """Cambia entre interpretación en TWT y en profundidad (en profundidad no hay capas)."""
        twt = self.v_modo.get() == "TWT"
        self.v_unid.set("s (TWT)" if twt else "km")
        for w in self.f_capas.winfo_children():
            self._estado(w, "normal" if twt else "disabled")
        self._rehacer_botones()
        self._dibujar()

    def _estado(self, w, estado):
        """Habilita o deshabilita un control y todos los que contiene."""
        try:
            w.configure(state=estado)
        except tk.TclError:          # algunos contenedores no tienen estado
            pass
        for c in w.winfo_children():
            self._estado(c, estado)

    # ------------------------------------------------------------------ capas
    def _dialogo_capas(self):
        """Pregunta cuántas capas tiene la línea y cuáles son (1 = Q, 2 = N, 3 = P)."""
        top = tk.Toplevel(self)
        top.title("Capas de la línea sísmica")
        top.transient(self)
        ttk.Label(top, padding=8, wraplength=420, text=(
            "¿Cuántas capas tiene la línea sísmica?\n\n"
            "Capa 1 = Cuaternario (Q)\nCapa 2 = Neógeno (N)\nCapa 3 = Paleógeno (P)\n\n"
            "Marque cuáles están presentes. La base de la capa más profunda presente se toma "
            "como el límite con el basamento.")).pack(anchor="w")
        n = tk.IntVar(value=len(self.capas))
        marcas = {k: tk.BooleanVar(value=k in self.capas) for k in ORDEN_CAPAS}
        g = ttk.Frame(top, padding=8)
        g.pack(anchor="w")
        ttk.Label(g, text="Número de capas:").grid(row=0, column=0, sticky="w")

        def por_numero():            # al cambiar el número se proponen las capas superiores
            for i, k in enumerate(ORDEN_CAPAS):
                marcas[k].set(i < n.get())
        ttk.Spinbox(g, from_=1, to=3, textvariable=n, width=4, command=por_numero,
                    state="readonly").grid(row=0, column=1, sticky="w")
        for i, k in enumerate(ORDEN_CAPAS):
            num, nombre, _, _ = CAPAS[k]
            ttk.Checkbutton(g, text=f"Capa {num} – {nombre}", variable=marcas[k]).grid(
                row=i + 1, column=0, columnspan=2, sticky="w")

        def ok():                    # se exige coherencia entre el número y las casillas
            sel = [k for k in ORDEN_CAPAS if marcas[k].get()]
            if len(sel) != n.get():
                messagebox.showwarning(APP, f"Indicó {n.get()} capa(s) pero marcó {len(sel)}.", parent=top)
                return
            self.capas = sel
            self.capas_definidas = True
            self._texto_capas()
            self._rehacer_botones()
            self._dibujar()
            top.destroy()
        ttk.Button(top, text="Aceptar", command=ok).pack(pady=8)
        top.grab_set()
        self.wait_window(top)

    def _texto_capas(self):
        """Resumen de las capas presentes y ausentes que se muestra en el panel."""
        nombres = ", ".join(f"{CAPAS[k][0]} {CAPAS[k][1]}" for k in self.capas)
        ausentes = [CAPAS[k][1] for k in ORDEN_CAPAS if k not in self.capas]
        txt = f"{len(self.capas)} capa(s): {nombres}."
        if ausentes:
            txt += f" Ausentes (espesor 0): {', '.join(ausentes)}."
        txt += f" Base de {CAPAS[self.capas[-1]][1]} = basamento."
        self.v_capas_txt.set(txt)

    def _horizontes_activos(self):
        """Horizontes a digitalizar: superficie, base de cada capa excepto la más profunda,
        y el límite con el basamento (que es la base de la capa más profunda presente)."""
        lista = [("sup", "Superficie / tope (opcional)", COLOR_SUP)]
        if self.v_modo.get() == "TWT":
            for k in self.capas[:-1]:
                num, nombre, _, color = CAPAS[k]
                lista.append((k, f"Base {nombre} (capa {num})", color))
            prof = CAPAS[self.capas[-1]][1]
            lista.append(("bas", f"Límite con basamento (= base {prof})", COLOR_BAS))
        else:
            lista.append(("bas", "Límite con basamento", COLOR_BAS))
        return lista

    def _rehacer_botones(self):
        """Regenera los botones de digitalización según el modo y las capas elegidas."""
        for w in self.f_hor.winfo_children():
            w.destroy()
        self._lbl_hor = {}
        for i, (k, texto, color) in enumerate(self._horizontes_activos()):
            tk.Label(self.f_hor, bg=color, width=2).grid(row=i, column=0, padx=2, pady=1)
            ttk.Button(self.f_hor, text=texto, width=34,
                       command=lambda k=k: self._activar(k)).grid(row=i, column=1, sticky="w")
            lbl = ttk.Label(self.f_hor, text=f"{len(self.hor[k])} pts")   # vértices marcados
            lbl.grid(row=i, column=2, padx=3)
            self._lbl_hor[k] = lbl

    def _nombre_hor(self, k):
        """Nombre legible de un horizonte."""
        return next((t for kk, t, _ in self._horizontes_activos() if kk == k), k)

    def _color_hor(self, k):
        """Color con que se dibuja un horizonte."""
        return {"sup": COLOR_SUP, "bas": COLOR_BAS}.get(k) or CAPAS[k][3]

    # ------------------------------------------------------------------ digitalización
    def _activar(self, modo):
        """Indica qué se marcará con los próximos clics (calibración u horizonte)."""
        if self.img is None:
            messagebox.showwarning(APP, "Primero cargue la imagen de la sección.")
            return
        self.activo = modo
        if modo == "xcal":
            self.xcal = []
            self.v_msg.set("ESCALA HORIZONTAL — Clic 1: extremo IZQUIERDO de la línea (km 0). "
                           "Clic 2: extremo DERECHO (km L). Clic derecho: deshacer.")
        elif modo == "ycal":
            self.ycal = []
            u = self.v_unid.get()
            self.v_msg.set(f"ESCALA VERTICAL — Clic 1: nivel de la referencia 1 ({self.v_ref1.get()} {u}). "
                           f"Clic 2: nivel de la referencia 2 ({self.v_ref2.get()} {u}).")
        else:
            self.v_msg.set(f"DIGITALIZANDO: {self._nombre_hor(modo)}. Clic izquierdo agrega vértices, "
                           "clic derecho deshace, doble clic o Enter termina. Rueda: zoom.")
        self._dibujar()
        self.canvas.get_tk_widget().focus_set()     # para que el lienzo reciba el teclado

    def _terminar(self):
        """Cierra la digitalización del horizonte activo."""
        if self.activo in self.hor:
            n = len(self.hor[self.activo])
            self.v_msg.set(f"{self._nombre_hor(self.activo)}: {n} vértices. Elija el siguiente horizonte.")
        self.activo = None
        self._dibujar()

    def _deshacer(self):
        """Elimina el último punto marcado en el modo activo."""
        if self.activo == "xcal" and self.xcal:
            self.xcal.pop()
        elif self.activo == "ycal" and self.ycal:
            self.ycal.pop()
        elif self.activo in self.hor and self.hor[self.activo]:
            self.hor[self.activo].pop()
        self._dibujar()

    def _borrar_horizonte(self):
        """Borra por completo el horizonte activo, previa confirmación."""
        if self.activo in self.hor and self.hor[self.activo]:
            if messagebox.askyesno(APP, f"¿Borrar {self._nombre_hor(self.activo)}?"):
                self.hor[self.activo] = []
                self._dibujar()
        else:
            messagebox.showinfo(APP, "Active primero el horizonte que desea borrar.")

    def _modo_toolbar(self):
        """Modo de la barra de matplotlib ('' si no está activo el zoom ni el desplazamiento)."""
        m = self.toolbar.mode
        return str(getattr(m, "value", m) or "")

    def _clic(self, ev):
        """Atiende los clics sobre la imagen según el modo activo."""
        # Se ignoran clics fuera de la imagen, sin modo activo o con zoom/desplazamiento activos.
        if ev.inaxes != self.ax or self.activo is None or self._modo_toolbar():
            return
        self.canvas.get_tk_widget().focus_set()
        if ev.button == 3:                      # botón derecho: deshacer
            self._deshacer()
            return
        if ev.button != 1:
            return
        if ev.dblclick and self.activo in self.hor:   # doble clic: terminar el horizonte
            self._terminar()
            return
        x, y = ev.xdata, ev.ydata               # posición en píxeles de la imagen
        if self.activo == "xcal":
            self.xcal.append(x)
            if len(self.xcal) == 2:
                self.xcal.sort()                # el menor es siempre el extremo izquierdo
                self.activo = None
                self.v_msg.set("Extremos marcados. Paso 3: marque las dos referencias verticales.")
        elif self.activo == "ycal":
            self.ycal.append(y)
            if len(self.ycal) == 2:
                self.activo = None
                self.v_msg.set("Escala vertical marcada. Defina las capas (TWT) y digitalice los horizontes.")
        else:
            if self.v_snap.get():
                x, y = self._iman(x, y)         # se pega a un horizonte cercano, si lo hay
            self.hor[self.activo].append([float(x), float(y)])
        self._dibujar()

    def _iman(self, x, y):
        """Proyecta el clic sobre el horizonte más cercano si está a menos de SNAP_PX píxeles.

        La distancia se mide en píxeles de pantalla, de modo que el imán se comporta igual
        con cualquier nivel de zoom. Se calcula, para cada segmento AB de los demás
        horizontes, el punto más cercano C = A + t·AB con t acotado a [0, 1].
        """
        p = self.ax.transData.transform((x, y))          # clic en coordenadas de pantalla
        mejor = (SNAP_PX, x, y)
        for k, pts in self.hor.items():
            if k == self.activo or len(pts) < 2:
                continue
            P = np.asarray(sorted(pts), float)
            D = self.ax.transData.transform(P)           # vértices en pantalla
            A, B = D[:-1], D[1:]
            AB = B - A
            t = np.clip(np.einsum("ij,ij->i", p - A, AB) / np.maximum(np.einsum("ij,ij->i", AB, AB), 1e-12), 0, 1)
            d = np.hypot(*(A + t[:, None] * AB - p).T)   # distancia a cada segmento
            i = int(np.argmin(d))
            if d[i] < mejor[0]:
                q = P[i] + t[i] * (P[i + 1] - P[i])      # punto equivalente en la imagen
                mejor = (d[i], q[0], q[1])
        return mejor[1], mejor[2]

    def _tecla(self, ev):
        """Atajos de teclado: Enter termina, Retroceso/Ctrl+Z deshace, Esc cancela."""
        if ev.key == "enter":
            self._terminar()
        elif ev.key in ("backspace", "ctrl+z"):
            self._deshacer()
        elif ev.key == "escape":
            self.activo = None
            self._dibujar()

    def _rueda(self, ev):
        """Zoom con la rueda del ratón centrado en el puntero (factor 1,25 por paso)."""
        if ev.inaxes != self.ax:
            return
        f = 1 / 1.25 if ev.button == "up" else 1.25
        x0, x1 = self.ax.get_xlim()
        y0, y1 = self.ax.get_ylim()
        self.ax.set_xlim(ev.xdata + (x0 - ev.xdata) * f, ev.xdata + (x1 - ev.xdata) * f)
        self.ax.set_ylim(ev.ydata + (y0 - ev.ydata) * f, ev.ydata + (y1 - ev.ydata) * f)
        self.canvas.draw_idle()

    def _dibujar(self):
        """Redibuja sobre la imagen la calibración, los horizontes y los pseudopozos."""
        for a in self._artistas:                 # se borra lo dibujado anteriormente
            try:
                a.remove()
            except (ValueError, NotImplementedError):
                pass
        self._artistas = []
        for k, lbl in getattr(self, "_lbl_hor", {}).items():
            lbl.configure(text=f"{len(self.hor[k])} pts")
        if self.img is None:
            self.canvas.draw_idle()
            return
        ax, A = self.ax, self._artistas.append
        H, W = self.img.shape[:2]
        caja = dict(boxstyle="round,pad=0.15", fc="white", ec="none", alpha=0.8)
        for i, x in enumerate(self.xcal):        # extremos de la sección (magenta)
            A(ax.axvline(x, color="magenta", ls="--", lw=1.2))
            A(ax.text(x, H * 0.97, "km 0" if i == 0 else f"km {self.v_long.get()}", color="magenta",
                      ha="center", fontsize=8, bbox=caja))
        refs = [self.v_ref1.get(), self.v_ref2.get()]
        for i, y in enumerate(self.ycal):        # referencias verticales (verde)
            A(ax.axhline(y, color="#00a000", ls="--", lw=1.2))
            A(ax.text(W * 0.005, y, f"{refs[i]} {self.v_unid.get()}", color="#00a000", va="bottom",
                      fontsize=8, bbox=caja))
        visibles = {k for k, _, _ in self._horizontes_activos()}
        for k, pts in self.hor.items():          # horizontes con sus vértices
            if not pts or k not in visibles:
                continue
            P = np.asarray(sorted(pts), float)
            c = self._color_hor(k)
            lw = 2.8 if k == self.activo else 1.8     # el horizonte activo se resalta
            A(ax.plot(P[:, 0], P[:, 1], "-", color=c, lw=lw)[0])
            A(ax.plot(P[:, 0], P[:, 1], "o", color=c, ms=3.5, mec="black", mew=0.4)[0])
        if self.pozos_km and len(self.xcal) == 2:     # verticales numeradas de los pseudopozos
            try:
                L = float(self.v_long.get())
                ini = int(self.v_punto_ini.get())
            except ValueError:
                L, ini = None, 1
            if L:
                for n, km in enumerate(self.pozos_km):
                    px = self.xcal[0] + km / L * (self.xcal[1] - self.xcal[0])
                    A(ax.axvline(px, color=COLOR_POZO, lw=0.9, alpha=0.85))
                    A(ax.text(px, H * 0.02, str(ini + n), color=COLOR_POZO, ha="center", va="top",
                              fontsize=7, fontweight="bold", bbox=caja))
        self.canvas.draw_idle()

    # ------------------------------------------------------------------ pseudopozos
    def _float(self, var, nombre):
        """Convierte el texto de una casilla en número (acepta coma decimal)."""
        try:
            return float(var.get().replace(",", "."))
        except ValueError:
            raise ValueError(f"Valor no válido en '{nombre}': {var.get()!r}")

    def _generar_pozos(self, silencioso=False):
        """Calcula las posiciones (km) de los pseudopozos por cantidad o por espaciado."""
        try:
            L = self._float(self.v_long, "Longitud de la línea")
            if L <= 0:
                raise ValueError("La longitud de la línea debe ser mayor que 0.")
            a = self._float(self.v_km_ini, "Desde km")
            b = self._float(self.v_km_fin, "Hasta km") if self.v_km_fin.get().strip() else L
            a, b = max(0.0, a), min(L, b)                  # el rango se acota a la línea
            if b <= a:
                raise ValueError("El rango de pseudopozos es vacío.")
            v = self._float(self.v_val_pozos, "Cantidad/Espaciado")
            if self.v_tipo_pozos.get() == "cantidad":
                n = int(round(v))
                if n < 1:
                    raise ValueError("La cantidad de pseudopozos debe ser al menos 1.")
                self.pozos_km = [a] if n == 1 else list(np.linspace(a, b, n))   # equiespaciados
            else:
                if v <= 0:
                    raise ValueError("El espaciado debe ser mayor que 0.")
                self.pozos_km = list(np.arange(a, b + 1e-9, v))   # cada v km, incluido el final
            int(self.v_punto_ini.get())                    # valida el primer punto de control
        except ValueError as e:
            if not silencioso:
                messagebox.showerror(APP, str(e))
            raise
        self._dibujar()
        if not silencioso:
            self.v_msg.set(f"{len(self.pozos_km)} pseudopozos entre km {a:.2f} y {b:.2f}. "
                           "Paso 7: Calcular y exportar.")

    def _cfg(self):
        """Valida que todos los pasos estén completos y reúne los parámetros del cálculo."""
        if self.linea_base is None:
            raise ValueError("No hay línea BIP seleccionada (paso 1).")
        if self.img is None:
            raise ValueError("No hay imagen cargada (paso 1).")
        if len(self.xcal) != 2:
            raise ValueError("Falta marcar los extremos izquierdo y derecho (paso 2).")
        if len(self.ycal) != 2 or abs(self.ycal[1] - self.ycal[0]) < 1:
            raise ValueError("Falta marcar las dos referencias verticales (paso 3).")
        r1 = self._float(self.v_ref1, "Referencia 1")
        r2 = self._float(self.v_ref2, "Referencia 2")
        if r1 == r2:
            raise ValueError("Las dos referencias verticales deben tener valores distintos.")
        modo = self.v_modo.get()
        if modo == "TWT" and not self.capas_definidas:
            raise ValueError("Defina las capas de la línea (paso 4).")
        if len(self.hor["bas"]) < 2:
            raise ValueError("Digitalice el límite con basamento (paso 5).")
        self._generar_pozos(silencioso=True)
        return {
            "linea": self.linea_valor,
            "modo": modo,
            "xcal": list(self.xcal), "ycal": list(self.ycal), "ref": [r1, r2],
            "long_km": self._float(self.v_long, "Longitud"),
            "ajustar_bip": bool(self.v_ajustar.get()),
            "offset_km": self._float(self.v_offset, "km BIP del extremo izquierdo"),
            "capas": list(self.capas),
            "vel": {k: self._float(self.v_vel[k], f"Velocidad {k}") for k in ORDEN_CAPAS},
            "horizontes": {k: [list(p) for p in v] for k, v in self.hor.items()},
            "pozos_km": [float(k) for k in self.pozos_km],
            "punto_inicial": int(self.v_punto_ini.get()),
        }

    def _exportar(self):
        """Calcula los pseudopozos y genera Excel, CSV, imagen de control y proyecto."""
        try:
            cfg = self._cfg()
        except ValueError as e:
            messagebox.showerror(APP, str(e))
            return
        nombre = re.sub(r"[^\w\-]+", "_", cfg["linea"])       # nombre de archivo seguro
        opciones = dict(title="Guardar base de pseudopozos", defaultextension=".xlsx",
                        initialfile=f"Pseudopozos_{nombre}.xlsx", filetypes=[("Excel", "*.xlsx")])
        try:
            ruta = filedialog.asksaveasfilename(confirmoverwrite=False, **opciones)
        except tk.TclError:          # versiones de Tk sin la opción confirmoverwrite
            ruta = filedialog.asksaveasfilename(**opciones)
        if not ruta:
            return
        agregar = False
        if os.path.exists(ruta):     # base existente: agregar la línea o reemplazar el archivo
            r = messagebox.askyesnocancel(
                APP, "El archivo ya existe.\n\nSí: AGREGAR esta línea a la base existente "
                     "(la numeración de puntos de control continúa).\nNo: REEMPLAZAR el archivo.")
            if r is None:
                return
            agregar = r
            if agregar:
                ult = ultimo_punto(ruta)
                if ult is None:
                    messagebox.showerror(APP, "El archivo no fue creado por esta herramienta "
                                              "(no tiene hoja 'Pseudopozos').")
                    return
                if cfg["punto_inicial"] <= ult:     # la numeración continúa sin repetirse
                    cfg["punto_inicial"] = ult + 1
                    self.v_punto_ini.set(str(ult + 1))
                    self._dibujar()
        linea = self.linea_base.orientada(self.v_extremo.get())   # km 0 = extremo izquierdo
        filas, avisos = calcular(cfg, linea)
        if not filas:
            messagebox.showerror(APP, "No se generó ningún pseudopozo.\n\n" + "\n".join(avisos[:15]))
            return
        meta = {                     # trazabilidad del procesamiento de esta línea
            "Fecha": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
            "Línea": cfg["linea"],
            "Tipo": "TWT" if cfg["modo"] == "TWT" else "Profundidad (km)",
            "Identificación": self.identificacion,
            "Archivo BIP": self.bip.ruta, "Capa BIP": self.bip.capa, "Campo nombre": self.linea_campo,
            "EPSG BIP": self.linea_base.epsg,
            "Desviación máx. respecto a línea BIP (m)": round(max(f["desv_m"] for f in filas), 6),
            "Imagen": self.img_ruta,
            "Longitud imagen (km)": cfg["long_km"],
            "Longitud BIP (km)": round(self.linea_base.longitud_km, 3),
            "Ubicación": "Proporcional a longitud BIP" if cfg["ajustar_bip"]
                         else f"km BIP = {cfg['offset_km']} + km imagen",
            "Extremo izquierdo": self.v_extremo.get(),
            "Capas presentes": ", ".join(cfg["capas"]) if cfg["modo"] == "TWT" else "NA",
            "Vel. Q (m/s)": cfg["vel"]["Q"], "Vel. N (m/s)": cfg["vel"]["N"], "Vel. P (m/s)": cfg["vel"]["P"],
            "Pseudopozos": len(filas),
            "Puntos de control": f"{filas[0]['punto']}–{filas[-1]['punto']}",
        }
        try:
            ruta_csv = exportar_excel(filas, meta, ruta, agregar)
        except PermissionError:      # el archivo está abierto en Excel
            messagebox.showerror(APP, "No se pudo escribir el archivo. Ciérrelo en Excel e intente de nuevo.")
            return
        except ValueError as e:
            messagebox.showerror(APP, str(e))
            return
        base = os.path.splitext(ruta)[0]
        ruta_png = f"{base}_{nombre}_control.png"
        ruta_json = f"{base}_{nombre}_proyecto.json"
        self._guardar_png(ruta_png)
        self._escribir_proyecto(ruta_json)
        esp = np.array([f["esp_total"] for f in filas])
        txt = (f"Se exportaron {len(filas)} pseudopozos de la línea {cfg['linea']}.\n\n"
               f"Espesor total: mín {esp.min():.1f} m, máx {esp.max():.1f} m, medio {esp.mean():.1f} m.\n"
               f"Desviación máxima respecto a la línea BIP ({self.linea_base.epsg}): "
               f"{max(f['desv_m'] for f in filas):.6f} m.\n\n"
               f"Excel: {ruta}\nCSV SIG: {ruta_csv}\nImagen de control: {ruta_png}\nProyecto: {ruta_json}")
        if avisos:
            txt += "\n\nAvisos:\n" + "\n".join(avisos[:12]) + ("\n…" if len(avisos) > 12 else "")
        messagebox.showinfo(APP, txt)
        self.v_msg.set(f"Exportado: {ruta}")

    def _guardar_png(self, ruta):
        """Guarda la sección completa con lo digitalizado como imagen de control de calidad."""
        xl, yl = self.ax.get_xlim(), self.ax.get_ylim()   # se recuerda el zoom actual
        H, W = self.img.shape[:2]
        self.ax.set_xlim(-0.5, W - 0.5)                   # encuadre de la imagen completa
        self.ax.set_ylim(H - 0.5, -0.5)
        dpi = 150
        tam = self.fig.get_size_inches()
        self.fig.set_size_inches(min(W, 4000) / dpi, min(W, 4000) / dpi * H / W)
        self.fig.savefig(ruta, dpi=dpi)
        self.fig.set_size_inches(*tam)                    # se restituye la vista del usuario
        self.ax.set_xlim(xl)
        self.ax.set_ylim(yl)
        self.canvas.draw_idle()

    # ------------------------------------------------------------------ proyecto
    def _escribir_proyecto(self, ruta):
        """Guarda en JSON todo lo necesario para retomar o auditar el trabajo de una línea."""
        datos = {
            "version": VERSION,
            "bip": {"ruta": self.bip.ruta if self.bip else None, "capa": self.bip.capa if self.bip else None,
                    "campo": self.linea_campo, "valor": self.linea_valor},
            "identificacion": self.identificacion,
            "imagen": self.img_ruta, "modo": self.v_modo.get(),
            "long_km": self.v_long.get(), "ajustar_bip": self.v_ajustar.get(),
            "offset_km": self.v_offset.get(), "extremo": self.v_extremo.get(),
            "xcal": self.xcal, "ycal": self.ycal, "ref": [self.v_ref1.get(), self.v_ref2.get()],
            "capas": self.capas, "capas_definidas": self.capas_definidas,
            "vel": {k: v.get() for k, v in self.v_vel.items()},
            "horizontes": self.hor,
            "pozos": {"tipo": self.v_tipo_pozos.get(), "valor": self.v_val_pozos.get(),
                      "ini": self.v_km_ini.get(), "fin": self.v_km_fin.get(),
                      "punto_ini": self.v_punto_ini.get()},
        }
        with open(ruta, "w", encoding="utf-8") as fh:
            json.dump(datos, fh, ensure_ascii=False, indent=1)

    def _guardar_proyecto(self):
        """Guarda el proyecto en la ruta que elija el usuario."""
        nombre = re.sub(r"[^\w\-]+", "_", self.linea_valor or "linea")
        ruta = filedialog.asksaveasfilename(title="Guardar proyecto", defaultextension=".json",
                                            initialfile=f"Proyecto_{nombre}.json",
                                            filetypes=[("Proyecto", "*.json")])
        if ruta:
            self._escribir_proyecto(ruta)
            self.v_msg.set(f"Proyecto guardado: {ruta}")

    def _abrir_proyecto(self):
        """Restituye un proyecto guardado sin volver a identificar ni preguntar la orientación."""
        ruta = filedialog.askopenfilename(title="Abrir proyecto", filetypes=[("Proyecto", "*.json")])
        if not ruta:
            return
        with open(ruta, encoding="utf-8") as fh:
            d = json.load(fh)
        if d.get("imagen") and os.path.exists(d["imagen"]):
            self._cargar_imagen(d["imagen"], identificar=False)
        elif d.get("imagen"):
            messagebox.showwarning(APP, f"No se encontró la imagen:\n{d['imagen']}")
        b = d.get("bip", {})
        self.linea_base = None
        if b.get("ruta") and os.path.exists(b["ruta"]):
            if self.bip is None or os.path.abspath(self.bip.ruta) != os.path.abspath(b["ruta"]):
                self._cargar_bip(b["ruta"], b.get("capa"), b.get("campo"), auto_identificar=False)
            if b.get("valor"):
                self._fijar_linea(b.get("campo"), b["valor"],
                                  d.get("identificacion") or "proyecto guardado", orientar=False)
        elif b.get("ruta"):
            messagebox.showwarning(APP, f"No se encontró el catálogo BIP del proyecto:\n{b['ruta']}")
        self.v_modo.set(d.get("modo", "TWT"))
        self.v_long.set(d.get("long_km", ""))
        self.v_ajustar.set(d.get("ajustar_bip", True))
        self.v_offset.set(d.get("offset_km", "0"))
        self.v_extremo.set(d.get("extremo", EXTREMOS[0]))
        self.v_ref1.set(d["ref"][0])
        self.v_ref2.set(d["ref"][1])
        self.capas = d.get("capas", ["Q", "N", "P"])
        self.capas_definidas = d.get("capas_definidas", False)
        for k, v in d.get("vel", {}).items():
            self.v_vel[k].set(v)
        p = d.get("pozos", {})
        self.v_tipo_pozos.set(p.get("tipo", "cantidad"))
        self.v_val_pozos.set(p.get("valor", "34"))
        self.v_km_ini.set(p.get("ini", "0"))
        self.v_km_fin.set(p.get("fin", ""))
        self.v_punto_ini.set(p.get("punto_ini", "1"))
        self.xcal, self.ycal = d.get("xcal", []), d.get("ycal", [])
        self.hor = {k: d.get("horizontes", {}).get(k, []) for k in self.hor}
        if self.capas_definidas:
            self._texto_capas()
        self._cambio_modo()
        try:
            self._generar_pozos(silencioso=True)
        except ValueError:
            pass
        self.v_msg.set(f"Proyecto abierto: {ruta}")


# =========================================================================================
# 12. PUNTO DE ENTRADA
# =========================================================================================
def main():
    """Inicia la aplicación. Con --autotest verifica PROJ, GDAL y OCR y se cierra sola."""
    try:
        import ctypes
        ctypes.windll.shcore.SetProcessDpiAwareness(1)   # nitidez en pantallas de alta resolución
    except Exception:
        pass
    app = App()
    if "--autotest" in sys.argv:
        import tempfile
        from PIL import ImageDraw, ImageFont
        try:
            e, n = Transformer.from_crs(3116, 9377, always_xy=True).transform(1000000.0, 1000000.0)
            prueba = os.path.join(tempfile.gettempdir(), "pseudopozos_ocr.png")
            im = Image.new("RGB", (700, 120), "white")
            ImageDraw.Draw(im).text((20, 30), "CV-1979-08", fill="black", font=ImageFont.load_default(48))
            im.save(prueba)
            ocr = ocr_imagen(prueba)
            estado = (f"OK proj={e:.1f},{n:.1f} gdal={pyogrio.__gdal_version_string__} "
                      f"drivers={len(pyogrio.list_drivers(read=True))} "
                      f"ocr={[t for t, _ in ocr] if ocr is not None else 'NO DISPONIBLE'}")
        except Exception as exc:
            estado = f"FALLO {exc!r}"
        with open(os.path.join(tempfile.gettempdir(), "pseudopozos_autotest.txt"), "w") as fh:
            fh.write(estado)
        app.after(4000, app.destroy)
    app.mainloop()


if __name__ == "__main__":      # solo se ejecuta al lanzar el archivo, no al importarlo
    main()
